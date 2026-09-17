"""Mẫu nhập trống theo cấu hình thật; không truy vấn dữ liệu khách hàng."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from core.constants import IMPORT_MAX_ROWS
from forms_builder import record_policies, choice_registry


def build(table):
    columns = list(table.columns.filter(is_computed=False).order_by('order', 'id'))
    policy = record_policies.for_table(table)
    if policy:
        columns += [c for c in policy.extra_columns(table) if not c.is_computed]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Nhap du lieu'
    guide = wb.create_sheet('Huong dan')
    choices = wb.create_sheet('Lua chon')
    choices.sheet_state = 'hidden'
    guide.append(['Hướng dẫn nhập vào ' + table.name, 'Nội dung'])
    guide.append(['Cách dùng', 'Điền từ hàng 2 của sheet Nhap du lieu, giữ nguyên tiêu đề. Mỗi hàng là một bản ghi.'])
    guide.append(['Nhập tệp', 'Mở đúng bảng → Nhập Excel → xem trước → xác nhận. Kiểm số dòng thành công và dòng lỗi sau khi xử lý.'])
    guide.append(['Giới hạn', f'Tối đa {IMPORT_MAX_ROWS} dòng mỗi lần. Mẫu không chứa dữ liệu khách hàng.'])
    guide.append(['Ô bắt buộc', 'Tiêu đề vàng là bắt buộc; ô không bắt buộc có thể để trống.'])
    guide.append(['Ngày / tiền', 'Ngày: yyyy-mm-dd. Tiền: số, không kèm ký hiệu tiền tệ. Giữ đúng loại tiền của đơn.'])
    guide.append(['Dòng mẫu', 'Ví dụ chỉ nằm ở sheet hướng dẫn, không được nhập tự động.'])
    for index, column in enumerate(columns, 1):
        letter = get_column_letter(index)
        cell = ws.cell(1, index, column.name)
        # Tên cột/lựa chọn do người dùng đặt luôn là chữ, không phải công thức.
        cell.data_type = 's'
        cell.font = Font(name='Calibri', size=11, bold=True, color='173D2C')
        cell.fill = PatternFill('solid', fgColor='FFF0BA' if column.required else 'DFEBDD')
        cell.alignment = Alignment(horizontal='left', wrap_text=True, vertical='center')
        ws.column_dimensions[letter].width = min(42, max(20, len(column.name) + 3))
        fmt = {'date': 'yyyy-mm-dd', 'datetime': 'yyyy-mm-dd hh:mm',
               'integer': '0', 'decimal': '0.00', 'money': '#,##0.00'}.get(column.field_type, '@')
        ws.column_dimensions[letter].number_format = fmt
        body_alignment = Alignment(wrap_text=True, vertical='top', horizontal='right' if column.field_type in ('integer', 'decimal', 'money') else 'left')
        ws.column_dimensions[letter].alignment = body_alignment
        ws.column_dimensions[letter].font = Font(name='Calibri', size=11)
        ws.cell(2, index).number_format = fmt
        ws.cell(2, index).alignment = body_alignment
        note = 'Bắt buộc.' if column.required else 'Có thể để trống.'
        if column.is_computed:
            note = 'Để trống: hệ thống tự tính, không nhận giá trị nhập cho cột này.'
        if policy and column.code in getattr(getattr(policy, 'assignment_service', None), 'COLUMNS', ()):
            note = 'Để trống. Phân công bằng hộp Phân công sau khi nhập.'
        if policy and column.code in getattr(policy, 'PROTECTED', ()):
            note = 'Nên để trống: hệ thống tính từ Chi tiết sản phẩm (JSON).'
        options = list(choice_registry.for_column(column).options()) if column.field_type == 'choice' else []
        if column.field_type == 'choice' and not column.is_computed:
            if options:
                for row, value in enumerate(options, 1):
                    option = choices.cell(row, index, str(value))
                    option.data_type = 's'
                name = f'LuaChon_{index}'
                wb.defined_names.add(DefinedName(name, attr_text=f"'Lua chon'!${letter}$1:${letter}${len(options)}"))
                validation = DataValidation(type='list', formula1=f'={name}', allow_blank=not column.required)
                validation.errorTitle = 'Giá trị không hợp lệ'
                validation.error = 'Chọn giá trị trong danh sách của cột.'
            else:
                note += ' Cột chưa có danh sách chọn: để trống; nếu bắt buộc, quản lý phải cấu hình trước khi nhập.'
                validation = DataValidation(type='custom', formula1=f'LEN({letter}2)=0', allow_blank=True)
                validation.error = 'Cột chưa có danh sách chọn. Quản lý cần cấu hình trước khi điền.'
            validation.showErrorMessage = True
            validation.errorStyle = 'stop'
            ws.add_data_validation(validation)
            validation.add(f'{letter}2:{letter}{IMPORT_MAX_ROWS + 1}')
        if column.code == getattr(policy, 'DETAIL_CODE', None):
            note = 'Bắt buộc. Điền chi tiết sản phẩm theo ví dụ trong sheet Huong dan; mã sản phẩm phải có trong danh mục.'
            guide.append(['Ví dụ Chi tiết sản phẩm (JSON)', '[{"product":"MA-SAN-PHAM","quantity":1,"unit_price":"100.00","paid_amount":"0.00"}]'])
            guide.append(['Chi tiết sản phẩm', 'Thay MA-SAN-PHAM bằng mã thật trong danh mục. Nhiều sản phẩm: thêm phần tử trong cùng danh sách.'])
        cell.comment = Comment(note, 'KN JSC')
        guide.append([column.name, note])
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 54
    ws.row_dimensions[2].height = 24
    ws.sheet_view.zoomScale = 85
    ws.auto_filter.ref = f'A1:{get_column_letter(len(columns))}1'
    guide.column_dimensions['A'].width = 38
    guide.column_dimensions['B'].width = 105
    for row in guide:
        for cell in row:
            cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        guide.row_dimensions[row[0].row].height = 44
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
