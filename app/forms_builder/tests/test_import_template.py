"""Mẫu tải xuống phải dùng được với bộ nhập thật, không chứa khách hàng."""
from datetime import date
from io import BytesIO

import pytest
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from forms_builder.models import TableDef, GrantAction
from forms_builder.services import import_service, grant_service
from forms_builder import record_policies
from orders.models import Product

pytestmark = pytest.mark.django_db


@pytest.mark.parametrize('code', ['van_don'])
def test_download_template_roundtrip(client, nguoi_dung, code, settings):
    settings.ROOT_URLCONF = 'knjsc.urls_bangtinh'
    call_command('tao_bang_van_don')
    table = TableDef.objects.get(code=code)
    client.force_login(nguoi_dung['admin'])
    response = client.get(f'/bang/{code}/mau-nhap.xlsx')
    assert response.status_code == 200
    assert 'attachment' in response['Content-Disposition']
    wb = load_workbook(BytesIO(response.content))
    ws = wb.worksheets[0]
    assert all(c.value is None for row in ws.iter_rows(min_row=2) for c in row)
    headers = [c.value for c in ws[1]]
    assert len(headers) == len(set(headers))
    columns = list(table.columns.order_by('order', 'id'))
    assert headers[:len(columns)] == [c.name for c in columns]
    values = {'ngay': date(2026, 9, 15), 'ten_khach': 'Khach kiem thu',
              'so_dien_thoai': '0012345678', 'ma_don': 'MAU-TEST-001', 'loai_tien': 'CAD', 'quoc_gia': 'Canada'}
    for i, column in enumerate(columns, 1):
        if column.code in values:
            ws.cell(2, i, values[column.code])
    policy = record_policies.for_table(table)
    if policy and hasattr(policy, 'DETAIL_CODE'):
        Product.objects.create(code='mau-sp', name='San pham test')
        index = headers.index('Chi tiết sản phẩm (JSON)') + 1
        ws.cell(2, index, '[{"product":"mau-sp","quantity":1,"unit_price":"100.00","paid_amount":"0.00"}]')
    for cell in ws[2]:
        if cell.value is not None:
            ws.cell(3, cell.column, cell.value)
    order_index = next(i for i, c in enumerate(columns, 1) if c.code == 'ma_don')
    ws.cell(3, order_index, 'MAU-TEST-002')
    stream = BytesIO()
    wb.save(stream)
    upload = SimpleUploadedFile('mau.xlsx', stream.getvalue())
    job = import_service.prepare(table, upload, actor=nguoi_dung['admin'])
    assert job.total == 2
    assert job.summary['ignored'] == []
    assert job.summary['preview_error_count'] == 0
    import_service.confirm(job, actor=nguoi_dung['admin'])
    job.refresh_from_db()
    assert job.summary['error_count'] == 0
    assert table.records.count() == 2
    phone_index = next(i for i, c in enumerate(columns, 1) if c.code == 'so_dien_thoai')
    assert ws.cell(2, phone_index).number_format == '@'
    assert job.summary['sample'][0][phone_index - 1] == '0012345678'
    listing = client.get('/thu-muc/?bp=van-don&tat-ca=1')
    assert listing.status_code == 200
    assert f'/bang/{code}/mau-nhap.xlsx' in listing.content.decode()
    assert '>Sửa</span>' not in listing.content.decode()


def test_template_choices_and_empty_choice_warning(client, nguoi_dung):
    call_command('tao_bang_van_don')
    table = TableDef.objects.get(code='van_don')
    from forms_builder.meaning import FieldType
    from forms_builder.models import ColumnDef
    column = ColumnDef.objects.create(table=table, code='kenh_thu', name='Kênh thử', field_type=FieldType.CHOICE, options=[], order=200)
    client.force_login(nguoi_dung['admin'])
    wb = load_workbook(BytesIO(client.get('/bang/van_don/mau-nhap.xlsx').content))
    ws = wb.worksheets[0]
    cell = next(c for c in ws[1] if c.value == column.name)
    assert 'chưa có danh sách chọn' in cell.comment.text
    currency = table.columns.get(code='loai_tien')
    index = next(c.column for c in ws[1] if c.value == currency.name)
    name = f'LuaChon_{index}'
    assert name in wb.defined_names
    assert [row[index - 1].value for row in wb['Lua chon'] if row[index - 1].value] == currency.options


def test_template_requires_import_permission(client, nguoi_dung, settings):
    settings.ROOT_URLCONF = 'knjsc.urls_bangtinh'
    call_command('tao_bang_van_don')
    staff = nguoi_dung['staff_sale_1']
    table = TableDef.objects.get(code='van_don')
    client.force_login(staff)
    assert client.get('/bang/van_don/mau-nhap.xlsx').status_code in (403, 404)
    grant_service.grant(table=table, user=staff, action=GrantAction.VIEW, actor=nguoi_dung['admin'])
    client.force_login(staff)
    assert client.get('/bang/van_don/mau-nhap.xlsx').status_code == 403
    assert '/bang/van_don/mau-nhap.xlsx' not in client.get('/thu-muc/?bp=van-don&tat-ca=1').content.decode()
    grant_service.grant(table=table, user=staff, action=GrantAction.EDIT, actor=nguoi_dung['admin'])
    client.force_login(staff)
    assert client.get('/bang/van_don/mau-nhap.xlsx').status_code == 200
    client.force_login(nguoi_dung['admin'])
    assert client.post('/bang/van_don/mau-nhap.xlsx').status_code == 405
