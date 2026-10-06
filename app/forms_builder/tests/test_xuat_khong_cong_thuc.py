"""Tệp Excel xuất ra không chạy công thức do người dùng gõ — AC-7.14 (săn lỗi bảo mật 06.10.2026).

openpyxl ghi mọi chuỗi bắt đầu bằng "=" thành **công thức**. Tên khách, ghi chú, địa chỉ là chữ do Sale gõ hay nhập
từ tệp: ai gõ `=HYPERLINK("http://…","Bấm")` thì người mở tệp xuất ra thấy một liên kết chạy được (chèn công thức
vào bảng tính). Nay mọi ô chữ ghi đúng là chữ, giữ nguyên nội dung, ở cả xuất trực tiếp, xuất nền (tệp lớn), Báo cáo
tổng hợp và tệp mẫu nhập.
"""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from core import excel
from forms_builder.services import import_template_service, record_service
from forms_builder.tests.test_nhap_xuat import bang_sale  # noqa: F401
from reports.tests.test_aggregations import bang_mkt  # noqa: F401
from reports.tests.test_mkt_excel import marketing_scope  # noqa: F401

pytestmark = pytest.mark.django_db

DOC = '=HYPERLINK("http://ke-xau.test","Bấm vào đây")'


def _o(wb_bytes_or_wb):
    wb = wb_bytes_or_wb if hasattr(wb_bytes_or_wb, "worksheets") else load_workbook(BytesIO(wb_bytes_or_wb))
    if not hasattr(wb_bytes_or_wb, "worksheets"):
        return [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    buf = BytesIO()
    wb.save(buf)
    return _o(buf.getvalue())


@pytest.mark.parametrize("ghi_lien", [False, True])
def test_write_table_ghi_chu_bat_dau_bang_dau_bang_la_chu(ghi_lien):
    """AC-7.14 — Ô chữ bắt đầu bằng "=" ghi thành chữ đúng nguyên văn, không thành công thức, ở cả chế độ thường và ghi liền (tệp lớn); số, ngày giữ kiểu"""
    o = _o(excel.write_table(["Tên", "Số"], [[DOC, 5], ["=1+1", 7]], write_only=ghi_lien))
    assert [c.data_type for c in o] == ["s", "s", "s", "n", "s", "n"]
    assert [c.value for c in o][2] == DOC


def test_xuat_bang_du_lieu_khong_co_cong_thuc(client, bang_sale, nguoi_dung):  # noqa: F811
    """AC-7.14 — Xuất Bảng dữ liệu có ô Khách hàng do người dùng gõ "=HYPERLINK(…)": tệp tải về không có ô công thức nào, ô đó vẫn đúng chữ đã gõ"""
    record_service.create_record(bang_sale, {"ngay": "2026-08-01", "khach": DOC, "doanh_thu": "100", "so_luong": 1},
                                 actor=nguoi_dung["manager_sale"])
    client.force_login(nguoi_dung["manager_sale"])
    o = _o(client.get("/bang/don_sale/xuat/").content)
    assert not [c.value for c in o if c.data_type == "f"]
    assert DOC in [c.value for c in o]


def test_tep_mau_nhap_khong_co_cong_thuc(bang_sale):  # noqa: F811
    """AC-7.14 — Tệp mẫu nhập: tên cột do quản lý đặt bắt đầu bằng "=" vẫn là chữ, không thành công thức"""
    bang_sale.columns.filter(code="khach").update(name="=1+1")
    o = _o(import_template_service.build(bang_sale))
    assert not [c.value for c in o if c.data_type == "f"] and "=1+1" in [c.value for c in o]


def test_xuat_bao_cao_tong_hop_khong_co_cong_thuc(client, marketing_scope, nguoi_dung):  # noqa: F811
    """AC-7.14 — Xuất Báo cáo tổng hợp: tên bảng do quản lý đặt bắt đầu bằng "=" vẫn là chữ trong tệp, không ô công thức nào"""
    bang = marketing_scope.table
    bang.name = DOC
    bang.save(update_fields=["name"])
    client.force_login(nguoi_dung["manager_sale"])
    kq = client.get("/bao-cao/tong-hop/xuat/", {"nguon": bang.code, "tu": "2026-08-01", "den": "2026-08-31"})
    assert kq.status_code == 200
    o = _o(kq.content)
    assert not [c.value for c in o if c.data_type == "f"]
    assert any(str(c.value).startswith(DOC) for c in o)
