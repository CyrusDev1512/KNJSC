"""Cột Ghi chú của bảng Vận đơn rộng gấp 2,5 lần mặc định và mang cờ tự giãn dòng — `docs/04` AC-11.44.

Ghi chú dài bị cắt ở cột 160 px, người dùng phải bấm mở hộp đọc mới xem được nội
dung. Bộ phận Vận đơn báo bất tiện (23.09.2026). Cột rộng ra và cờ `auto_height`
là vế máy chủ; vế trình duyệt — dòng tự giãn cao vừa nội dung, kể cả ngắt dòng
người dùng gõ — nằm ở bài `tests/e2e/test_ghi_chu_tu_gian_dong.py`.

Chỉ bảng vận đơn đổi: lưới là một bộ dùng chung cho mọi bảng (ADR-021) và **không
nhận diện nghiệp vụ bằng mã cột** — cờ do profile bảng trả trong metadata, nên bài
này kiểm luôn rằng bảng khác có cột `ghi_chu` vẫn 160 px và không có cờ.

Hai chỗ kề bên cũng thuộc AC-11.44: ô Ghi chú ở Lên đơn là ô nhiều dòng (Sale gõ
được xuống dòng), và tệp Excel xuất ra bật Wrap Text cho cột văn bản dài.
"""
import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.http import QueryDict
from openpyxl import load_workbook

from crm.services.master_grid_service import metadata
from crm.tests.test_waybill_feedback import feedback  # noqa: F401  (fixture dùng lại)
from forms_builder.meaning import FieldType, Meaning
from forms_builder.models import DataRecord, TableDef
from forms_builder.services import export_service, import_service
from orders.constants import ACTIVE_WAYBILL_TABLE_CODE, Market, PaymentMethod
from orders.models import Order
from orders.services import dispatch_service

pytestmark = pytest.mark.django_db

#: Mặc định của mọi cột trong `master_grid_service.metadata`
RONG_MAC_DINH = 160
#: Gấp 2,5 lần — đủ cho một câu ghi chú không phải xuống dòng ngay
RONG_GHI_CHU = 400


def _cot(bang, ma):
    return next(c for c in metadata(list(bang.columns.order_by("order"))) if c["code"] == ma)


def _bang_van_don(nguoi_dung):
    dispatch_service.ensure_waybill_table(actor=nguoi_dung["admin"])
    return TableDef.objects.get(code=ACTIVE_WAYBILL_TABLE_CODE)


def _bang_khac(departments, nguoi_dung):
    """Bảng Sale thường cũng có cột văn bản dài mang mã `ghi_chu` — không phải vận đơn."""
    bang = TableDef.objects.create(
        name="Đơn hàng Sale", code="don_sale_ghi_chu",
        department=departments["sale"], created_by=nguoi_dung["manager_sale"],
    )
    bang.columns.create(name="Ngày", code="ngay", field_type=FieldType.DATE,
                        meaning=Meaning.DATE, order=0)
    bang.columns.create(name="Ghi chú", code="ghi_chu",
                        field_type=FieldType.LONG_TEXT, order=1)
    return bang


# ── Độ rộng cột ──────────────────────────────────────────────────────────────

def test_cot_ghi_chu_van_don_rong_gap_hai_ruoi(nguoi_dung):
    """AC-11.44 — Cột Ghi chú của bảng Vận đơn rộng 400 px mặc định."""
    assert _cot(_bang_van_don(nguoi_dung), "ghi_chu")["width"] == RONG_GHI_CHU


def test_cot_khac_cua_van_don_giu_nguyen(nguoi_dung):
    """AC-11.44 — Chỉ Ghi chú rộng ra; cột khác của bảng Vận đơn vẫn 160 px."""
    bang = _bang_van_don(nguoi_dung)
    for ma in ("ten_khach", "so_dien_thoai", "dia_chi"):
        assert _cot(bang, ma)["width"] == RONG_MAC_DINH, ma


def test_bang_khac_khong_bi_anh_huong(departments, nguoi_dung):
    """AC-11.44 — Bảng không phải vận đơn: cột ghi chú vẫn 160 px (ADR-021)."""
    assert _cot(_bang_khac(departments, nguoi_dung), "ghi_chu")["width"] == RONG_MAC_DINH


# ── Cờ tự giãn dòng: profile bảng quyết định, không phải mã cột ─────────────

def test_co_tu_gian_dong_chi_o_ghi_chu_van_don(nguoi_dung):
    """AC-11.44 — Profile Vận đơn gắn cờ `auto_height` cho riêng cột Ghi chú; cột khác của bảng không có cờ."""
    bang = _bang_van_don(nguoi_dung)
    assert _cot(bang, "ghi_chu")["auto_height"] is True
    for ma in ("ten_khach", "so_dien_thoai", "dia_chi", "bill"):
        assert not _cot(bang, ma).get("auto_height"), ma


def test_bang_khac_khong_co_co_tu_gian(departments, nguoi_dung):
    """AC-11.44 — Bảng không phải vận đơn: cột văn bản dài `ghi_chu` không mang cờ — lưới không nhận diện theo mã cột (ADR-021)."""
    assert "auto_height" not in _cot(_bang_khac(departments, nguoi_dung), "ghi_chu")


def test_co_tu_gian_di_theo_json_khoi(client, feedback, nguoi_dung):  # noqa: F811
    """AC-11.44 — JSON `du-lieu/` của lưới mang cờ `auto_height` và rộng 400 tới trình duyệt."""
    bang = feedback[0]
    client.force_login(nguoi_dung["staff_vd"])
    kq = client.get(f"/bang-tinh/{bang.code}/du-lieu/")
    assert kq.status_code == 200
    cot = {c["code"]: c for c in kq.json()["columns"]}
    assert cot["ghi_chu"]["auto_height"] is True
    assert cot["ghi_chu"]["width"] == RONG_GHI_CHU
    assert not cot["ten_khach"].get("auto_height")


# ── Hai chỗ kề bên làm mất ngắt dòng của ghi chú ─────────────────────────────

def test_len_don_ghi_chu_o_nhieu_dong(client, feedback, nguoi_dung):  # noqa: F811
    """AC-11.44 — Lên đơn: ô Ghi chú là ô nhiều dòng; xuống dòng trình duyệt gửi (CRLF) đi vào `ghi_chu` thành `\n` gọn."""
    client.force_login(nguoi_dung["staff_sale_1"])
    trang = client.get("/van-don/len-don/")
    assert trang.status_code == 200
    assert '<textarea name="note"' in trang.content.decode()

    data = {
        "customer_name": "Khách hai dòng", "phone": "0912345678",
        "market": Market.US, "payment_method": PaymentMethod.ZELLE,
        # Trình duyệt gửi textarea theo CRLF; hệ thống chỉ giữ một kiểu xuống dòng
        "note": "Giao sau 17h\r\nGọi trước 15 phút",
        "product": [feedback[1][0].code], "quantity": ["1"], "unit_price": ["1"], "paid_amount": ["0"],
    }
    assert client.post("/van-don/len-don/", data).status_code == 200
    # Đơn lưu ghi chú đã gom về `\n`, và dòng vận đơn sinh ra mang y nguyên
    don = Order.objects.get(note="Giao sau 17h\nGọi trước 15 phút")
    assert don.record.data["ghi_chu"] == "Giao sau 17h\nGọi trước 15 phút"


def test_nhap_tep_excel_giu_ngat_dong(feedback, nguoi_dung):  # noqa: F811
    """AC-11.44 — Nhập tệp Excel có ô Ghi chú xuống dòng bằng Alt+Enter: ký tự xuống dòng đi
    vào dữ liệu nguyên vẹn, chỉ cắt khoảng trắng hai đầu — lưới mới giãn dòng đúng được."""
    bang, _, dong = feedback
    cot = list(bang.columns.all())
    ma_cot = [c.code for c in cot]
    ghi_chu = "  Giao buổi sáng\nGọi trước 30 phút\nKhách hay vắng  "

    # Xuất một dòng có sẵn ra rồi sửa ô Ghi chú: tệp nhập luôn khớp đúng cấu trúc bảng
    wb = export_service.build_workbook(DataRecord.objects.filter(pk=dong[0].pk), cot, title="VD")
    wb.active.cell(2, ma_cot.index("ma_don") + 1, "NHAP-XUONG-DONG")
    wb.active.cell(2, ma_cot.index("ghi_chu") + 1, ghi_chu)
    dem = io.BytesIO()
    wb.save(dem)

    viec = import_service.prepare(
        bang, SimpleUploadedFile("ghi-chu.xlsx", dem.getvalue()), actor=nguoi_dung["admin"])
    assert viec.summary["preview_error_count"] == 0, viec.summary
    import_service.confirm(viec, actor=nguoi_dung["admin"])

    moi = DataRecord.objects.filter(table=bang).order_by("-pk").first()
    assert moi.data["ma_don"] == "NHAP-XUONG-DONG", moi.data.get("ma_don")
    assert moi.data["ghi_chu"] == "Giao buổi sáng\nGọi trước 30 phút\nKhách hay vắng"


@pytest.mark.parametrize("ghi_lien", [False, True])
def test_xuat_excel_cot_ghi_chu_xuong_dong(feedback, nguoi_dung, settings, ghi_lien):  # noqa: F811
    """AC-11.44 — Tệp Excel xuất từ lưới Vận đơn bật Wrap Text ở cột Ghi chú, cột khác giữ nguyên — cả hai chế độ ghi."""
    settings.CRM_OPT_EXPORT = ghi_lien
    bang, _, dong = feedback
    dong[0].data["ghi_chu"] = "Dòng 1\nDòng 2"
    dong[0].save(update_fields=["data"])

    kind, wb = export_service.export(nguoi_dung["admin"], bang, QueryDict("trang=10"), builder="grid")
    assert kind == "file"
    dem = io.BytesIO()
    wb.save(dem)
    dem.seek(0)
    ws = load_workbook(dem).active
    tieu_de = [c.value for c in ws[1]]
    cot = tieu_de.index("Ghi chú") + 1
    o = ws.cell(row=2, column=cot)
    assert o.alignment.wrap_text is True and o.alignment.vertical == "top"
    assert "\n" in (o.value or "") or "\n" in (ws.cell(row=3, column=cot).value or "")
    # Cột đầu bảng không phải văn bản dài — không bị bật Wrap Text theo
    assert not ws.cell(row=2, column=1).alignment.wrap_text
