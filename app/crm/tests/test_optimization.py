"""AC-24.1–24.4, AC-24.6: hồi quy tối ưu trên fixture, không dùng DB local."""
import uuid
import pytest
from django.db import connection, transaction
from django.test.utils import CaptureQueriesContext
from .test_waybill_feedback import feedback, delivery_leader
from .test_master_grid import BASE, write

pytestmark = pytest.mark.django_db(transaction=True)


@pytest.fixture
def enabled(settings):
    from django.core.cache import caches
    settings.CACHES={'default':{'BACKEND':'django.core.cache.backends.locmem.LocMemCache'},'crm':{'BACKEND':'django.core.cache.backends.locmem.LocMemCache','LOCATION':'optimization-tests'}}
    caches['crm'].clear()
    settings.CRM_OPT_READ = True
    settings.CRM_OPT_SYNC = True
    settings.CRM_OPT_RECEIPTS = True
    settings.CRM_OPT_STATS = True


def test_v2_repeat_block_has_no_count(client, feedback, nguoi_dung, enabled):
    client.force_login(nguoi_dung['admin'])
    first = client.get(BASE+'du-lieu/', {'protocol': 2}).json()
    assert first['protocol'] == 2
    with CaptureQueriesContext(connection) as queries:
        second = client.get(BASE+'du-lieu/', {'protocol':2,'query_token':first['query_token']}).json()
    assert second['total'] == first['total']
    assert not any('COUNT(' in q['sql'].upper() for q in queries)


def test_v2_value_change_keeps_query_and_rollback_does_not_publish(client, feedback, nguoi_dung, enabled):
    client.force_login(nguoi_dung['admin']);row=feedback[2][0]
    first=client.get(BASE+'du-lieu/',{'protocol':2}).json()
    assert first['protocol']==2
    result=write(client,row,column='ghi_chu',old=row.data.get('ghi_chu'),value='Kiểm phiên bản')
    assert result.status_code==200
    current=client.get(BASE+'du-lieu/',{'protocol':2}).json()
    assert current['query_token']==first['query_token']
    assert current['revision']>first['revision']
    with transaction.atomic():
        row.refresh_from_db();row.data['ghi_chu']='Không commit';row.save()
        transaction.set_rollback(True)
    assert client.get(BASE+'du-lieu/',{'protocol':2}).json()['revision']==current['revision']


def test_v2_receipt_is_compact_and_replay_does_not_duplicate_history(client, feedback, nguoi_dung, enabled):
    from crm.models import GridCellHistory, GridMutationReceipt
    client.force_login(nguoi_dung['admin']);row=feedback[2][0]
    payload={'protocol':2,'operation':str(uuid.uuid4()),'cells':[{'id':row.pk,'column':'ghi_chu','old':row.data.get('ghi_chu'),'value':'Gọn'}]}
    first=client.post(BASE+'luu-json/',payload,content_type='application/json')
    assert first.status_code==200
    assert first.json()['protocol']==2
    assert 'rows' not in first.json()
    count=GridCellHistory.objects.count()
    replay=client.post(BASE+'luu-json/',payload,content_type='application/json').json()
    assert replay['cells']==first.json()['cells'] and replay['replayed']
    assert GridCellHistory.objects.count()==count
    assert 'rows' not in GridMutationReceipt.objects.get(operation=payload['operation']).result


def test_compact_receipt_survives_flag_rollback(client,feedback,nguoi_dung,enabled,settings):
    client.force_login(nguoi_dung['admin']);row=feedback[2][0]
    payload={'protocol':2,'operation':str(uuid.uuid4()),'cells':[{'id':row.pk,'column':'ghi_chu','old':row.data.get('ghi_chu'),'value':'A'}]}
    original=client.post(BASE+'luu-json/',payload,content_type='application/json').json()
    settings.CRM_OPT_RECEIPTS=False
    assert write(client,row,column='ghi_chu',old='A',value='B').status_code==200
    replay=client.post(BASE+'luu-json/',payload,content_type='application/json')
    assert replay.status_code==200
    assert replay.json()=={**original,'replayed':True}


def test_sync_only_changed_visible_rows_and_no_scope_count(client,feedback,nguoi_dung,enabled):
    client.force_login(nguoi_dung['admin']);rows=feedback[2]
    first=client.get(BASE+'du-lieu/',{'protocol':2}).json()
    assert write(client,rows[0],column='ghi_chu',old=rows[0].data.get('ghi_chu'),value='Đổi ô').status_code==200
    payload={'query_token':first['query_token'],'revision':first['revision'],'ids':[r.pk for r in rows],'visible':[rows[0].pk]}
    with CaptureQueriesContext(connection) as queries:
        response=client.post(BASE+'dong-bo/',payload,content_type='application/json')
    assert response.status_code==200
    assert [r['id'] for r in response.json()['rows']]==[rows[0].pk]
    assert not response.json()['reset']
    assert not any('COUNT(' in q['sql'].upper() or 'MAX(' in q['sql'].upper() for q in queries)


def test_warm_cache_loses_assignment_immediately(client,feedback,nguoi_dung,delivery_leader,enabled):
    from .test_waybill_feedback import assign_rows
    row=feedback[2][0];staff=nguoi_dung['staff_vd']
    assign_rows(delivery_leader,[row],delivery=staff.pk)
    client.force_login(staff)
    before=client.get(BASE+'du-lieu/',{'protocol':2}).json()
    assert before['total']==1
    assign_rows(delivery_leader,[row],delivery=None)
    sync=client.post(BASE+'dong-bo/',{'query_token':before['query_token'],'revision':before['revision'],'ids':[row.pk],'visible':[row.pk]},content_type='application/json').json()
    assert sync['removed']==[row.pk] and sync['reset'] and not sync['rows']
    assert client.get(BASE+'du-lieu/',{'protocol':2}).json()['total']==0
    assert client.get(BASE+'du-lieu/',{'protocol':2,'query_token':before['query_token']}).status_code==409


def test_bulk_and_nested_writes_publish_one_commit(feedback):
    from crm.models import GridRevision,GridChange,GridPendingChange
    from forms_builder.models import DataRecord
    table,_,rows=feedback
    before=GridRevision.objects.get(table=table).revision
    with transaction.atomic():
        DataRecord.objects.filter(pk=rows[0].pk).update(style={'ghi_chu':{'fs':16}})
        with transaction.atomic():DataRecord.objects.filter(pk=rows[1].pk).update(style={'ghi_chu':{'fs':20}})
        assert GridRevision.objects.get(table=table).revision==before
    current=GridRevision.objects.get(table=table).revision
    assert current==before+1
    event=GridChange.objects.get(table=table,revision=current)
    assert set(event.record_ids)=={r.pk for r in rows}
    assert event.columns==['__style'] and not GridPendingChange.objects.exists()


def test_streamed_excel_preserves_values_types_and_header():
    from core.excel import write_table
    from openpyxl import load_workbook
    from io import BytesIO
    from decimal import Decimal
    from datetime import date
    rows=[['00123',Decimal('123.45'),date(2026,9,11),None,'Tiếng Việt\nxuống dòng']]
    snapshots=[]
    for streaming in (False,True):
        wb=write_table(['Mã','Tiền','Ngày','Rỗng','Chữ'],iter(rows),write_only=streaming)
        out=BytesIO();wb.save(out);out.seek(0)
        sheet=load_workbook(out).active
        snapshots.append((sheet.freeze_panes,[(c.coordinate,c.value,c.data_type,c.number_format,c.font.bold) for row in sheet for c in row]))
    assert snapshots[0]==snapshots[1]


def test_statistics_cache_expiry_refresh_and_access(feedback,nguoi_dung,enabled):
    from crm.services import statistics_cache
    from django.http import QueryDict
    from django.utils import timezone
    from unittest.mock import Mock,patch
    from datetime import timedelta
    from .test_waybill_feedback import assign_rows
    table=feedback[0];user=nguoi_dung['admin'];now=timezone.now();build=Mock(return_value={'orders':2})
    def get(params=''):
        return statistics_cache.result(user,[table],QueryDict(params),now.date(),now.date(),build)
    with patch.object(statistics_cache.timezone,'now',return_value=now):first=get()
    with patch.object(statistics_cache.timezone,'now',return_value=now+timedelta(seconds=10)):
        assert get()==first and build.call_count==1
    with patch.object(statistics_cache.timezone,'now',return_value=now+timedelta(seconds=16)):
        assert get()['calculated']!=first['calculated'] and build.call_count==2
    get('lam_moi=1');assert build.call_count==3
    assign_rows(user,[feedback[2][0]],delivery=nguoi_dung['staff_vd'].pk)
    get();assert build.call_count==4


def test_redis_failure_reads_authoritative_data(client,feedback,nguoi_dung,enabled):
    from unittest.mock import patch
    client.force_login(nguoi_dung['admin'])
    with patch('django.core.cache.backends.locmem.LocMemCache.get',side_effect=ConnectionError):
        result=client.get(BASE+'du-lieu/',{'protocol':2})
    assert result.status_code==200 and result.json()['total']==len(feedback[2])


def test_statistics_concurrent_miss_reuses_first_calculation(feedback,nguoi_dung,enabled):
    """AC-24.6: hai lần mở cùng cache key chỉ tính một lần khi lease còn hiệu lực."""
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    from unittest.mock import patch
    from django.core.cache import caches
    from django.db import connections
    from django.http import QueryDict
    from django.utils import timezone
    from crm.services import statistics_cache
    user=nguoi_dung['admin'];table=feedback[0];day=timezone.localdate()
    user.profile  # Đọc quan hệ trên thread tạo fixture trước khi chia đối tượng chỉ đọc.
    building=Event();waiting=Event();release=Event();calls=[]
    cache_type=type(caches['crm']);original=cache_type.add
    def add(cache,key,*args,**kwargs):
        value=original(cache,key,*args,**kwargs)
        if key.endswith(':computing') and not value:waiting.set()
        return value
    def build():
        calls.append(1);building.set()
        assert release.wait(5)
        return {'orders':2,'groups':[1,2]}
    def get():
        try:return statistics_cache.result(user,[table],QueryDict(''),day,day,build)
        finally:connections.close_all()
    with patch.object(cache_type,'add',add),ThreadPoolExecutor(max_workers=2) as pool:
        first=pool.submit(get)
        try:
            assert building.wait(5)
            second=pool.submit(get)
            assert waiting.wait(5)
        finally:release.set()
        a,b=first.result(timeout=5),second.result(timeout=5)
    assert len(calls)==1 and a==b
    a['value']['groups'].append(3)
    assert b['value']['groups']==[1,2]


def test_signed_cursor_forward_backward_and_scope_binding(client,feedback,nguoi_dung,enabled):
    from forms_builder.models import DataRecord
    from django.utils import timezone
    from datetime import timedelta
    table,_,rows=feedback
    DataRecord.objects.bulk_create([DataRecord(table=table,data={'ma_don':f'C-{i}'},created_by=nguoi_dung['admin'],created_at=timezone.now()+timedelta(seconds=i)) for i in range(220)])
    client.force_login(nguoi_dung['admin'])
    first=client.get(BASE+'du-lieu/',{'protocol':2}).json()
    params={'protocol':2,'offset':100,'query_token':first['query_token'],'cursor':first['next_cursor']}
    second=client.get(BASE+'du-lieu/',params).json()
    offset=client.get(BASE+'du-lieu/',{'protocol':2,'offset':100}).json()
    assert [r['id'] for r in second['rows']]==[r['id'] for r in offset['rows']]
    back=client.get(BASE+'du-lieu/',{'protocol':2,'offset':0,'cursor':second['previous_cursor']}).json()
    assert [r['id'] for r in back['rows']]==[r['id'] for r in first['rows']]
    assert client.get(BASE+'du-lieu/',{**params,'cursor':params['cursor']+'bad'}).status_code==409
    client.force_login(nguoi_dung['staff_sale_1'])
    assert client.get(BASE+'du-lieu/',params).status_code==409


def test_revision_commit_order_not_transaction_start_order(feedback):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    from django.db import close_old_connections
    from forms_builder.models import DataRecord
    from crm.models import GridChange,GridRevision
    table,_,rows=feedback;before=GridRevision.objects.get(table=table).revision
    waiting=Event();later_committed=Event()
    def earlier():
        close_old_connections()
        try:
            with transaction.atomic():
                DataRecord.objects.filter(pk=rows[0].pk).update(style={'ghi_chu':{'fs':16}})
                waiting.set();assert later_committed.wait(10)
        finally:connection.close()
    def later():
        close_old_connections()
        try:
            assert waiting.wait(10)
            DataRecord.objects.filter(pk=rows[1].pk).update(style={'ghi_chu':{'fs':20}})
            later_committed.set()
        finally:connection.close()
    with ThreadPoolExecutor(max_workers=2) as pool:
        a=pool.submit(earlier);b=pool.submit(later);a.result(20);b.result(20)
    changes=list(GridChange.objects.filter(table=table,revision__gt=before).order_by('revision'))
    assert [e.record_ids for e in changes]==[[rows[1].pk],[rows[0].pk]]


def test_revision_migration_forward_backward_preserves_business(feedback):
    from django.db.migrations.executor import MigrationExecutor
    from forms_builder.models import DataRecord
    before=list(DataRecord.objects.order_by('id').values_list('id','data','style'))
    executor=MigrationExecutor(connection)
    try:
        executor.migrate([('crm','0002_grid_cell_history')])
        assert list(DataRecord.objects.order_by('id').values_list('id','data','style'))==before
        executor=MigrationExecutor(connection);executor.migrate([('crm','0005_assignment_journal_columns')])
        assert list(DataRecord.objects.order_by('id').values_list('id','data','style'))==before
    finally:
        MigrationExecutor(connection).migrate([('crm','0005_assignment_journal_columns')])


def test_journal_retention_and_large_statement_reset(feedback,nguoi_dung):
    from crm.models import GridRevision,GridChange
    from forms_builder.models import DataRecord
    table=feedback[0]
    GridChange.objects.filter(table=table).delete()
    GridRevision.objects.filter(table=table).update(revision=10001)
    GridChange.objects.bulk_create([GridChange(table=table,revision=i) for i in range(1,10002)],batch_size=500)
    DataRecord.objects.bulk_create([DataRecord(table=table,data={},created_by=nguoi_dung['admin']) for _ in range(2001)],batch_size=2500)
    event=GridChange.objects.filter(table=table).latest('revision')
    assert event.reset and event.record_ids==[]
    assert GridChange.objects.filter(table=table).count()==10000
    assert GridChange.objects.filter(table=table).earliest('revision').revision==3


def test_assignment_journal_uses_actual_column_codes(feedback,nguoi_dung):
    from crm.models import GridChange
    from .test_waybill_feedback import assign_rows
    from orders.services.assignment_service import COLUMNS
    assign_rows(nguoi_dung['admin'],[feedback[2][0]],delivery=nguoi_dung['staff_vd'].pk)
    event=GridChange.objects.filter(table=feedback[0]).latest('revision')
    assert set(COLUMNS).issubset(event.columns)


def test_streaming_export_direct_background_equivalence_and_revocation(feedback,nguoi_dung,settings,tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    from unittest.mock import patch
    from django.http import QueryDict
    from core.constants import JobStatus
    from core.exceptions import BusinessError
    from forms_builder.services import export_service
    from .test_waybill_feedback import assign_rows
    table,_,rows=feedback;user=nguoi_dung['staff_vd']
    assign_rows(nguoi_dung['admin'],rows,delivery=user.pk)
    settings.EXPORT_DIR=tmp_path/'exports';settings.STORAGE_DIR=tmp_path
    expected=None;jobs=[]
    for flag in (False,True):
        settings.CRM_OPT_EXPORT=flag
        _,book=export_service.export(user,table,QueryDict('sap=ma_don'),builder='grid')
        file=BytesIO();book.save(file);file.seek(0)
        values=list(load_workbook(file).active.values)
        if expected is None:expected=values
        assert values==expected
        with patch.object(export_service,'EXPORT_SYNC_MAX_ROWS',0),patch('forms_builder.services.import_service._day_vao_hang_doi'):
            _,job=export_service.export(user,table,QueryDict('sap=ma_don'),builder='grid')
        export_service.run(job.pk);job.refresh_from_db();jobs.append(job)
        assert job.status==JobStatus.DONE
        assert list(load_workbook(export_service.result_file(job)).active.values)==expected
        assert set(job.summary['exported_row_ids'])=={r.pk for r in rows}
    assign_rows(nguoi_dung['admin'],[rows[0]],delivery=None)
    for job in jobs:
        with pytest.raises(BusinessError,match='xuất lại'):export_service.check_download(job,user)
