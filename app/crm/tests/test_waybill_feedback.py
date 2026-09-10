"""Feedback KN CRM — AC-20.1, AC-20.2, AC-20.3, AC-20.4, AC-20.5, AC-20.7.

Quyền phân công, CAS, truy vấn dùng chung, Excel và migration.
"""
import pytest
from django.http import QueryDict

from forms_builder.models import DataRecord, TableDef
from orders.constants import ACTIVE_WAYBILL_TABLE_CODE
from orders.models import Product
from orders.services import dispatch_service, order_service
from crm.services import grid_service
from core.constants import Rank, JobStatus
from core.exceptions import BusinessError, OutOfScopeError
from forms_builder.models import GrantAction
from forms_builder.services import grant_service, record_service, export_service, import_service
from orders.models import WaybillAssignment, WaybillItem
from orders.services import assignment_service as assignment
from crm.services import sidebar_service
from django.utils import timezone
from datetime import timedelta
from unittest.mock import patch
from io import BytesIO
from django.core.files.uploadedfile import SimpleUploadedFile

pytestmark = pytest.mark.django_db


@pytest.fixture
def feedback(departments, nguoi_dung, settings):
    settings.GRID_ONLY_TABLES = set()
    dispatch_service.ensure_waybill_table(actor=nguoi_dung['admin'])
    table = TableDef.objects.get(code=ACTIVE_WAYBILL_TABLE_CODE)
    products = [Product.objects.create(code=f'feedback-{i}', name=f'Sản phẩm {i}') for i in range(2)]
    rows = []
    for i, product in enumerate(products):
        order = order_service.create_order(phone=f'090000000{i}', customer_name=f'Khách {i}',
            lines=[{'product': product.code, 'quantity': 1, 'unit_price': '10.00'}],
            actor=nguoi_dung['staff_sale_1' if i == 0 else 'staff_sale_2'])
        rows.append(order.record)
    return table, products, rows


def test_unassigned_rows_hidden_from_delivery_staff(feedback, nguoi_dung):
    assert not DataRecord.objects.in_scope(nguoi_dung['staff_vd']).filter(table=feedback[0]).exists()


def test_product_filter_matches_item_code(feedback, nguoi_dung):
    table, products, rows = feedback
    grid = grid_service.build_grid(nguoi_dung['admin'], QueryDict(f'sp={products[0].code}'), table=table)
    assert list(grid.queryset.values_list('pk', flat=True)) == [rows[0].pk]
    header = grid_service.build_grid(nguoi_dung['admin'], QueryDict(f'f_san_pham__trong={products[0].code}'), table=table)
    assert list(header.queryset.values_list('pk', flat=True)) == [rows[0].pk]


@pytest.fixture
def delivery_leader(make_user, departments):
    return make_user('vd_leader', Rank.LEADER, departments['vd'])


def assign_rows(user, rows, **changes):
    versions = {r.pk: WaybillAssignment.objects.filter(record=r).values_list('version', flat=True).first() or 0 for r in rows}
    return assignment.assign(user, versions, changes)


def visible(user, table):
    return set(DataRecord.objects.in_scope(user).filter(table=table).values_list('pk', flat=True))


def test_assignment_scope_and_view_only_care(feedback, nguoi_dung, delivery_leader, make_user, departments):
    """AC-20.1 — Quyền dòng theo tài khoản, giao CSKH chỉ bổ sung xem."""
    table, _, rows = feedback
    other = make_user('vd_other', Rank.STAFF, departments['vd'])
    from org.models import Department
    care = make_user('care_person', Rank.STAFF, Department.objects.create(code='cskh', name='CSKH'))
    assign_rows(delivery_leader, [rows[0]], delivery=nguoi_dung['staff_vd'].pk, care=care.pk, marketing=nguoi_dung['staff_mkt'].pk)
    assign_rows(delivery_leader, [rows[1]], delivery=other.pk, care=nguoi_dung['staff_sale_1'].pk)
    assert visible(nguoi_dung['staff_vd'], table) == {rows[0].pk}
    assert visible(other, table) == {rows[1].pk}
    assert visible(care, table) == {rows[0].pk}
    assert visible(nguoi_dung['staff_sale_1'], table) == {r.pk for r in rows}
    assert visible(nguoi_dung['staff_sale_2'], table) == {rows[1].pk}
    assert not visible(nguoi_dung['staff_mkt'], table)
    assert visible(delivery_leader, table) == {r.pk for r in rows}
    assert not grant_service.can_edit_record(care, rows[0])
    assert not grant_service.can_edit_record(nguoi_dung['staff_sale_1'], rows[1])
    grant_service.grant(table=table, user=other, action=GrantAction.EDIT, actor=nguoi_dung['admin'])
    assert visible(other, table) == {rows[1].pk}
    assert not grant_service.can_edit_record(other, rows[0])


def test_bulk_keep_clear_conflict_is_atomic(feedback, nguoi_dung, delivery_leader):
    """AC-20.3 — Giữ/bỏ/đổi và rollback toàn bộ lượt xung đột."""
    rows = feedback[2]
    assign_rows(delivery_leader, rows, delivery=nguoi_dung['staff_vd'].pk, care=nguoi_dung['staff_sale_1'].pk)
    assign_rows(delivery_leader, [rows[1]], marketing=nguoi_dung['staff_mkt'].pk)
    with pytest.raises(BusinessError, match='vừa được'):
        assignment.assign(delivery_leader, {r.pk: 1 for r in rows}, {'delivery': None})
    assert WaybillAssignment.objects.filter(delivery=nguoi_dung['staff_vd']).count() == 2
    assign_rows(delivery_leader, rows, delivery=None)
    assert not visible(nguoi_dung['staff_vd'], feedback[0])
    assert WaybillAssignment.objects.filter(care=nguoi_dung['staff_sale_1']).count() == 2
    assert WaybillAssignment.objects.get(record=rows[1]).marketing_id == nguoi_dung['staff_mkt'].pk


@pytest.mark.parametrize('invalid', ['department', 'inactive', 'locked'])
def test_assignment_invalid_account_all_or_nothing(feedback, nguoi_dung, delivery_leader, invalid):
    user = nguoi_dung['staff_sale_1'] if invalid == 'department' else nguoi_dung['staff_vd']
    if invalid == 'inactive':
        user.is_active = False; user.save(update_fields=['is_active'])
    if invalid == 'locked':
        user.profile.locked_until = timezone.now() + timedelta(hours=1)
        user.profile.save(update_fields=['locked_until'])
    with pytest.raises(BusinessError):
        assign_rows(delivery_leader, feedback[2], delivery=user.pk)
    assert not WaybillAssignment.objects.exists()


@pytest.mark.parametrize('role', ['staff_vd', 'staff_sale_1', 'manager_sale', 'manager_mkt'])
def test_assignment_denies_non_delivery_management(feedback, nguoi_dung, role, client):
    with pytest.raises(OutOfScopeError):
        assign_rows(nguoi_dung[role], feedback[2], delivery=nguoi_dung['staff_vd'].pk)
    client.force_login(nguoi_dung[role])
    assert client.get('/van-don/phan-cong/', {'row': feedback[2][0].pk}).status_code == 403


def test_assignment_endpoint_versions_and_audit(feedback, nguoi_dung, delivery_leader, client):
    from core.models import AuditLog
    client.force_login(delivery_leader)
    rows = feedback[2]
    response = client.get('/van-don/phan-cong/', {'row': [r.pk for r in rows]})
    assert response.status_code == 200
    assert all(r['version'] == 0 for r in response.json()['rows'])
    data = {'versions': {r.pk: 0 for r in rows}, 'changes': {'delivery': nguoi_dung['staff_vd'].pk}}
    assert client.post('/van-don/phan-cong/', data, content_type='application/json').status_code == 200
    assert client.post('/van-don/phan-cong/', data, content_type='application/json').status_code == 409
    audits = AuditLog.objects.filter(detail__startswith='Phân công')
    assert audits.count() == 2
    assert not audits.filter(detail__contains='09000000').exists()


@pytest.mark.parametrize('value', [True, 1.5, '1.5', '9' * 100])
def test_assignment_rejects_non_integer_account(feedback, delivery_leader, value):
    with pytest.raises(BusinessError):
        assign_rows(delivery_leader, feedback[2], delivery=value)


@pytest.mark.parametrize('operation', ['cell', 'paste', 'style', 'styles', 'delete'])
def test_old_assignee_cannot_write_stale_record(feedback, nguoi_dung, delivery_leader, operation):
    row = feedback[2][0]; user = nguoi_dung['staff_vd']
    assign_rows(delivery_leader, [row], delivery=user.pk)
    assign_rows(delivery_leader, [row], delivery=None)
    with pytest.raises(OutOfScopeError):
        if operation == 'cell': record_service.update_cell(row, 'ghi_chu', 'bad', actor=user)
        elif operation == 'paste': record_service.update_cells([(row, 'ghi_chu', 'bad')], actor=user)
        elif operation == 'style': record_service.update_style(row, 'ghi_chu', {'b': 1}, actor=user)
        elif operation == 'styles': record_service.update_styles([(row, 'ghi_chu')], {'b': 1}, actor=user)
        else: record_service.delete_record(row, actor=user)


@pytest.mark.parametrize('code', assignment.COLUMNS)
def test_assignment_columns_block_generic_writes(feedback, nguoi_dung, code):
    """AC-20.2 — Chỉ đường phân công được ghi người phụ trách."""
    row = feedback[2][0]
    with pytest.raises(BusinessError, match='Phân công'):
        record_service.update_cell(row, code, 'staff_vd', actor=nguoi_dung['admin'])
    with pytest.raises(BusinessError):
        record_service.update_cells([(row, 'ghi_chu', 'bad'), (row, code, 'bad')], actor=nguoi_dung['admin'])
    row.refresh_from_db()
    assert row.data.get('ghi_chu') != 'bad'


def test_filters_and_statistics_use_scoped_rows(feedback, nguoi_dung, delivery_leader, client):
    """AC-20.4 — Lọc và thống kê sau phạm vi, Marketing có chưa gán."""
    from orders.services import waybill_service
    table, products, rows = feedback
    assign_rows(delivery_leader, [rows[0]], delivery=nguoi_dung['staff_vd'].pk, marketing=nguoi_dung['staff_mkt'].pk)
    user = nguoi_dung['staff_vd']
    options = sidebar_service.product_options(user, table, list(table.columns.all()), QueryDict())
    assert [o[0] for o in options['items']] == [products[0].code]
    marketer_col = table.columns.get(code='phu_trach_mkt')
    assert grid_service.filter_options(user, table, marketer_col) == [('staff_mkt', 1)]
    params = QueryDict(f'sp={products[0].code}&f_phu_trach_mkt__trong=staff_mkt&f_quoc_gia__trong=US&trang=9')
    # Dùng giá trị thị trường thực tế của dòng, không suy đoán mã/nhãn.
    params = params.copy(); params['f_quoc_gia__trong'] = rows[0].data['quoc_gia']
    grid = grid_service.build_grid(user, params, table=table)
    assert set(grid.queryset.values_list('pk', flat=True)) == {rows[0].pk}
    stats, _ = waybill_service.statistics(grid.queryset)
    assert sum(r['orders'] for r in stats) == 1
    params['f_phu_trach_mkt__trong'] = '__unassigned__'
    assert not grid_service.build_grid(user, params, table=table).queryset.exists()
    assert grid_service.build_grid(nguoi_dung['admin'], QueryDict('f_phu_trach_mkt__trong=__unassigned__'), table=table).queryset.get().pk == rows[1].pk
    client.force_login(user)
    assert client.get(f'/van-don/chi-tiet/{rows[1].pk}/').status_code == 403
    assert client.get('/thong-ke/').status_code == 200


def test_product_or_without_duplicates_and_states_and(feedback, nguoi_dung):
    table, products, rows = feedback
    WaybillItem.objects.create(record=rows[0], product=products[1], quantity=1, unit_price='1.00')
    params = QueryDict(f'sp={products[0].code}&sp={products[1].code}')
    grid = grid_service.build_grid(nguoi_dung['admin'], params, table=table)
    assert grid.queryset.count() == 2
    params = params.copy()
    params['f_trang_thai_vc__trong'] = rows[0].data['trang_thai_vc']
    params['f_trang_thai_tt__trong'] = 'Thanh toán 1 phần'
    assert not grid_service.build_grid(nguoi_dung['admin'], params, table=table).queryset.exists()


def test_export_employee_codes_and_async_revoke(feedback, nguoi_dung, delivery_leader, settings, tmp_path, client):
    """AC-20.5 — Mã nhân viên và quyền file trực tiếp/nền."""
    table, _, rows = feedback
    assign_rows(delivery_leader, rows, delivery=nguoi_dung['staff_vd'].pk, care=nguoi_dung['staff_sale_2'].pk)
    user = nguoi_dung['staff_vd']
    for role in ['staff_sale_1', 'staff_sale_2']:
        profile = nguoi_dung[role].profile; profile.full_name = 'Trùng tên'; profile.save(update_fields=['full_name'])
    kind, wb = export_service.export(user, table, QueryDict('trang=10'), builder='grid')
    assert kind == 'file'
    values = list(wb.active.values)
    assert len(values) == 3
    headers = values[0]
    assert {r[headers.index('Mã Sale tạo đơn')] for r in values[1:]} == {'staff_sale_1', 'staff_sale_2'}
    assert {r[headers.index('Mã nhân viên Vận đơn')] for r in values[1:]} == {'staff_vd'}
    settings.STORAGE_DIR = tmp_path; settings.EXPORT_DIR = tmp_path / 'exports'
    with patch.object(export_service, 'EXPORT_SYNC_MAX_ROWS', 0), patch('forms_builder.services.import_service._day_vao_hang_doi'):
        _, job = export_service.export(user, table, QueryDict('trang=10'), builder='grid')
    export_service.run(job.pk); job.refresh_from_db()
    assert job.status == JobStatus.DONE
    from openpyxl import load_workbook
    direct = BytesIO(); wb.save(direct); direct.seek(0)
    assert list(load_workbook(export_service.result_file(job)).active.values) == list(load_workbook(direct).active.values)
    assert set(job.summary['exported_row_ids']) == {r.pk for r in rows}
    assign_rows(delivery_leader, [rows[0]], delivery=None)
    with pytest.raises(BusinessError, match='xuất lại'):
        export_service.check_download(job, user)
    client.force_login(user)
    from django.urls import reverse
    response = client.get(reverse('tac_vu_tai', args=[job.pk]))
    assert response.status_code == 302 and 'attachment' not in response.get('Content-Disposition', '')


def test_import_cannot_set_assignment(feedback, nguoi_dung):
    table, _, rows = feedback
    columns = list(table.columns.all())
    wb = export_service.build_workbook(DataRecord.objects.filter(pk=rows[0].pk), columns, title='VD')
    wb.active.cell(2, [c.code for c in columns].index('phu_trach_vd') + 1, 'staff_vd')
    stream = BytesIO(); wb.save(stream)
    job = import_service.prepare(table, SimpleUploadedFile('bad.xlsx', stream.getvalue()), actor=nguoi_dung['admin'])
    assert job.summary['preview_error_count'] == 1
    assert not WaybillAssignment.objects.exists()


def test_export_date_filter_and_unlinked_sale_blank(feedback, nguoi_dung):
    table, _, rows = feedback
    record_service.update_cell(rows[0], 'ngay', '2026-08-01', actor=nguoi_dung['admin'])
    record_service.update_cell(rows[1], 'ngay', '2026-08-02', actor=nguoi_dung['admin'])
    params = QueryDict('f_ngay__lon_bang=2026-08-01&f_ngay__nho_bang=2026-08-01&trang=99')
    _, wb = export_service.export(nguoi_dung['admin'], table, params, builder='grid')
    values = list(wb.active.values)
    assert len(values) == 2 and values[1][0] == rows[0].data['ma_don']
    imported = DataRecord.objects.create(table=table, department=table.department, data=rows[0].data)
    _, wb = export_service.export(nguoi_dung['admin'], table, params, builder='grid')
    values = list(wb.active.values)
    assert len(values) == 3
    codes = [r[values[0].index('Mã Sale tạo đơn')] for r in values[1:]]
    assert 'staff_sale_1' in codes and '' in codes


@pytest.mark.django_db(transaction=True)
def test_two_assigners_do_not_silently_overwrite(feedback, nguoi_dung, delivery_leader):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from django.db import connections
    from django.contrib.auth import get_user_model
    barrier = Barrier(2)
    row_ids = [r.pk for r in feedback[2]]
    leader_id = delivery_leader.pk
    target = nguoi_dung['staff_vd'].pk
    def attempt(reverse):
        try:
            user = get_user_model().objects.get(pk=leader_id)
            barrier.wait(timeout=10)
            assignment.assign(user, {pk: 0 for pk in (row_ids[::-1] if reverse else row_ids)}, {'delivery': target})
            return 'saved'
        except BusinessError as exc:
            return exc.code
        finally:
            connections.close_all()
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(attempt, [False, True]))
    assert sorted(results) == ['conflict', 'saved']
    assert set(WaybillAssignment.objects.values_list('version', flat=True)) == {1}


@pytest.mark.django_db(transaction=True)
def test_assignment_migration_roundtrip_preserves_records(feedback):
    """AC-20.7 — Chạy ngược/xuôi không gán hoặc đổi dữ liệu khách."""
    from django.db import connection
    from django.db.migrations.executor import MigrationExecutor
    table, _, rows = feedback
    expected = {r.pk: r.data for r in rows}
    try:
        MigrationExecutor(connection).migrate([('orders', '0003_waybillitem')])
        MigrationExecutor(connection).migrate([('orders', '0005_assignment_columns')])
        assert not WaybillAssignment.objects.exists()
        assert table.columns.filter(code__in=assignment.COLUMNS).count() == 3
        assert dict(DataRecord.objects.filter(pk__in=expected).values_list('pk', 'data')) == expected
    finally:
        MigrationExecutor(connection).migrate([('orders', '0005_assignment_columns')])


def test_download_worker_rechecks_current_assignments(feedback, nguoi_dung, delivery_leader, settings, tmp_path):
    table, _, rows = feedback
    user = nguoi_dung['staff_vd']
    assign_rows(delivery_leader, rows, delivery=user.pk)
    settings.STORAGE_DIR = tmp_path; settings.EXPORT_DIR = tmp_path / 'exports'
    with patch.object(export_service, 'EXPORT_SYNC_MAX_ROWS', 0), patch('forms_builder.services.import_service._day_vao_hang_doi'):
        _, job = export_service.export(user, table, QueryDict(), builder='grid')
    assign_rows(delivery_leader, [rows[0]], delivery=None)
    export_service.run(job.pk); job.refresh_from_db()
    assert job.status == JobStatus.DONE
    assert job.summary['exported_row_ids'] == [rows[1].pk]


def test_grid_ui_and_filtered_url(feedback, nguoi_dung, delivery_leader, client):
    client.force_login(delivery_leader)
    params = {'sp': feedback[1][0].code, 'f_trang_thai_tt__trong': 'Chưa thanh toán', 'sap': 'ma_don', 'trang': '2'}
    response = client.get('/bang-tinh/van_don_moi/', params)
    assert response.status_code == 200
    quick = response.context['quick_filters']
    assert ('sap', 'ma_don') in quick['keep'] and not any(k == 'trang' for k, _ in quick['keep'])
    html = response.content.decode()
    assert 'Phân công' in html and 'Thanh toán 1 phần' in html
    assert 'id="vd-entry"' not in html
    assert 'Bộ lọc' in html


def test_erp_reads_new_assignment_scope(feedback, nguoi_dung, client, settings):
    settings.ROOT_URLCONF = 'knjsc.urls'
    client.force_login(nguoi_dung['staff_vd'])
    response = client.get('/bang/van_don_moi/')
    assert response.status_code == 200
    assert 'Khách 0' not in response.content.decode() and 'Khách 1' not in response.content.decode()
    assert TableDef.objects.in_scope(nguoi_dung['staff_vd']).with_visible_record_count(nguoi_dung['staff_vd']).get(pk=feedback[0].pk).so_dong == 0


def test_scoped_grid_does_not_query_per_row(feedback, nguoi_dung, delivery_leader, client, django_assert_max_num_queries):
    table, _, rows = feedback
    more = DataRecord.objects.bulk_create([DataRecord(table=table, department=table.department,
        created_by=nguoi_dung['staff_sale_1'], data={**rows[0].data, 'ma_don': f'PERF-{i}'}) for i in range(98)])
    WaybillAssignment.objects.bulk_create([WaybillAssignment(record=r, delivery=nguoi_dung['staff_vd']) for r in rows + more])
    client.force_login(nguoi_dung['staff_vd'])
    client.get('/bang-tinh/van_don_moi/moi-nhat/')
    with django_assert_max_num_queries(22):
        response = client.get('/bang-tinh/van_don_moi/du-lieu/')
    assert response.status_code == 200 and len(response.json()['rows']) == 100
