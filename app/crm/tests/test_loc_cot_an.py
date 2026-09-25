"""Lọc theo cột đang ẩn vẫn chạy, kèm lời nhắc — bổ sung ADR-039, 22.09.2026 (TL-53).

Trước đây bộ lọc chỉ đọc trên cột đang hiện: ẩn cột đi là bộ lọc trỏ tới nó bị
**bỏ lặng lẽ**, lưới hiện thừa dòng mà không ai hiểu vì sao (URL cũ, liên kết
Thống kê, bookmark). Nay bộ lọc đọc trên mọi cột; chỗ nào có lọc theo cột ẩn thì
KN CRM hiện chip cảnh báo "(cột đang ẩn) …" bỏ được bằng nút ×, Bảng dữ liệu KN
ERP hiện dòng nhắc cạnh nút Xoá lọc. Cột ẩn vẫn không hiện ra ở đâu.
"""
from io import BytesIO

import pytest
from django.test import override_settings
from openpyxl import load_workbook

from forms_builder.services import export_service, table_service
from orders.services import dispatch_service

from .test_an_cot import luoi  # noqa: F401  (bảng vận đơn có cột sản phẩm)
from .test_waybill_feedback import feedback  # noqa: F401

pytestmark = pytest.mark.django_db


@pytest.fixture
def cot_an(luoi, nguoi_dung):  # noqa: F811
    """Bảng vận đơn có một cột sản phẩm đã ẩn, hai dòng số lượng 7 và 2."""
    bang, products, rows = luoi
    ma = dispatch_service.product_column_code(products[0])
    for dong, so_luong in zip(rows, (7, 2)):
        dong.data[ma] = so_luong
        dong.save(update_fields=["data"])
    table_service.set_columns_hidden(bang, [ma], True, actor=nguoi_dung["admin"])
    return bang, ma, rows


def test_loc_theo_cot_an_van_chay_va_chip_canh_bao(client, cot_an, nguoi_dung):
    """AC-39.8 — Bộ lọc trỏ tới cột đang ẩn vẫn lọc đúng dòng trên lưới KN CRM; chip mang
    nhãn "(cột đang ẩn)" và cờ cảnh báo, liên kết của chip bỏ đúng bộ lọc đó; bộ lọc
    trên cột đang hiện không bị dính cờ; cột ẩn vẫn không có trong danh sách cột trả về"""
    bang, ma, rows = cot_an
    client.force_login(nguoi_dung["staff_vd"])

    data = client.get(f"/bang-tinh/{bang.code}/du-lieu/?f_{ma}__lon_bang=5").json()
    assert [d["id"] for d in data["rows"]] == [rows[0].pk]
    assert ma not in [c["code"] for c in data["columns"]]

    trang = client.get(f"/bang-tinh/{bang.code}/?f_{ma}__lon_bang=5&f_ten_khach__chua=Kh")
    chips = {nhan: (url, an) for nhan, url, an in trang.context["chips"]}
    canh_bao = next(n for n in chips if n.startswith("(cột đang ẩn)"))
    assert "Sản phẩm 0" in canh_bao
    url, an = chips[canh_bao]
    assert an is True and f"f_{ma}__lon_bang" not in url and "f_ten_khach__chua" in url
    binh_thuong = next(n for n in chips if "Tên khách" in n)
    assert chips[binh_thuong][1] is False and not binh_thuong.startswith("(cột đang ẩn)")
    assert 'class="mg-chip-an"' in trang.content.decode()


def test_xuat_excel_theo_dung_bo_loc_cot_an(client, cot_an, nguoi_dung):
    """AC-39.8 — Tệp Excel xuất từ lưới theo đúng bộ lọc cột ẩn đang áp (ADR-002 "xuất
    đúng thứ đang hiện"), nhưng cột ẩn vẫn không có trong tiêu đề tệp"""
    bang, ma, rows = cot_an
    client.force_login(nguoi_dung["admin"])
    tep = client.get(f"/bang-tinh/{bang.code}/xuat/?f_{ma}__lon_bang=5")
    assert tep.status_code == 200
    trang_tinh = load_workbook(BytesIO(tep.content)).active
    hang = list(trang_tinh.iter_rows(values_only=True))
    assert len(hang) == 2  # tiêu đề + đúng một dòng khớp
    assert "Sản phẩm 0" not in hang[0]


@override_settings(ROOT_URLCONF="knjsc.urls")
def test_bang_du_lieu_erp_loc_cot_an_co_dong_nhac(client, cot_an, nguoi_dung):
    """AC-39.8 — Bảng dữ liệu KN ERP: lọc theo cột ẩn vẫn chạy và trang mang dòng nhắc
    "Đang lọc theo cột đang ẩn" kèm tên cột; không lọc cột ẩn thì không có dòng nhắc;
    bộ đọc chung của xuất tệp (`export_service.build_queryset`) cũng áp bộ lọc đó"""
    bang, ma, rows = cot_an
    client.force_login(nguoi_dung["admin"])

    kq = client.get(f"/bang/{bang.code}/?f_{ma}__lon_bang=5")
    assert kq.status_code == 200
    assert [d.pk for d, _ in kq.context["cac_dong"]] == [rows[0].pk]
    assert kq.context["loc_cot_an"] == ["Sản phẩm 0"]
    assert "Đang lọc theo cột đang ẩn" in kq.content.decode()

    sach = client.get(f"/bang/{bang.code}/")
    assert sach.context["loc_cot_an"] == [] and "Đang lọc theo cột đang ẩn" not in sach.content.decode()

    ds, cot, _ = export_service.build_queryset(
        nguoi_dung["admin"], bang, {"f_" + ma + "__lon_bang": "5"})
    assert list(ds.values_list("pk", flat=True)) == [rows[0].pk]
    assert ma not in [c.code for c in cot]


@override_settings(ROOT_URLCONF="knjsc.urls")
def test_loc_cot_an_khong_mo_rong_pham_vi(client, cot_an, nguoi_dung):
    """AC-39.8 — Chiều bị từ chối: lọc theo cột ẩn không mở đường xem dòng ngoài phạm vi —
    nhân viên Sale vẫn chỉ thấy dòng của mình dù bộ lọc cột ẩn khớp cả dòng người khác"""
    bang, ma, rows = cot_an
    client.force_login(nguoi_dung["staff_sale_2"])
    assert client.get(f"/bang/{bang.code}/?f_{ma}__lon_bang=1").status_code == 403
    with override_settings(ROOT_URLCONF="knjsc.urls_bangtinh"):
        kq = client.get(f"/bang-tinh/{bang.code}/du-lieu/?f_{ma}__lon_bang=1")
    assert kq.status_code == 200
    thay = [d["id"] for d in kq.json()["rows"]]
    assert rows[0].pk not in thay  # dòng của staff_sale_1 khớp bộ lọc nhưng ngoài phạm vi
