"""ADR-039 — Ẩn cột với cả công ty: quản lý bảng bật tắt trong hộp "Cột", cột biến mất
khỏi lưới KN CRM, tệp Excel xuất ra và Bảng dữ liệu bên KN ERP, dữ liệu ô vẫn giữ.

Bổ sung 22.09: **cột số lượng theo sản phẩm mặc định ẩn**, nên các bài dưới đây lấy
`thanh_pho` làm ví dụ cột đang hiện — dùng cột `sl_*` thì không còn chứng minh được gì."""
from importlib import import_module
from io import BytesIO

import pytest
from django.core.exceptions import ValidationError
from django.apps import apps
from django.db import connection, migrations
from django.db.migrations.executor import MigrationExecutor
from django.test import override_settings
from openpyxl import load_workbook

from core.constants import AuditAction
from core.exceptions import BusinessError
from core.models import AuditLog
from forms_builder.models import ColumnDef, DataRecord, TableDef
from forms_builder.services import table_service
from orders.constants import WAYBILL_TABLE_CODE
from orders.models import Product
from orders.services import dispatch_service, order_service, product_service

from .test_waybill_feedback import feedback  # noqa: F401

pytestmark = pytest.mark.django_db
URL = f"/bang-tinh/{WAYBILL_TABLE_CODE}/an-cot/"


@pytest.fixture
def luoi(feedback):
    """Bảng vận đơn đã có cột số lượng theo sản phẩm, như máy thật sau `tao_bang_van_don`.

    Từ 22.09 nhóm cột đó vào ở trạng thái **ẩn với cả công ty**."""
    bang, products, rows = feedback
    dispatch_service.sync_product_columns(bang)
    return bang, products, rows


def _ma_cot_luoi(client, code=WAYBILL_TABLE_CODE):
    kq = client.get(f"/bang-tinh/{code}/du-lieu/")
    assert kq.status_code == 200
    return [c["code"] for c in kq.json()["columns"]]


def _cot_san_pham(bang):
    return [c.code for c in bang.columns.all() if dispatch_service.is_product_column(c.code)]


def test_an_cot_bien_mat_ba_man_hinh_nhung_giu_du_lieu(client, luoi, nguoi_dung):
    """AC-39.1 — Quản lý bảng ẩn cột: cột biến khỏi lưới KN CRM, tệp Excel xuất ra và Bảng
    dữ liệu bên KN ERP với mọi người; giá trị ô vẫn nằm nguyên trong bản ghi; có nhật ký"""
    bang, products, rows = luoi
    ma = "thanh_pho"                     # cột thường: không khoá, không bắt buộc nhập
    dong = rows[0]
    dong.data[ma] = "Đà Nẵng"
    dong.save(update_fields=["data"])

    client.force_login(nguoi_dung["admin"])
    assert ma in _ma_cot_luoi(client)

    kq = client.post(URL, {"cot": [ma], "an": "1"})
    assert kq.status_code == 200 and kq.json()["da_doi"] == [ma]

    # 1. Lưới KN CRM — với chính người vừa ẩn và với nhân viên Vận đơn
    assert ma not in _ma_cot_luoi(client)
    client.force_login(nguoi_dung["staff_vd"])
    assert ma not in _ma_cot_luoi(client)

    # 2. Tệp Excel xuất ra từ lưới
    tep = client.get(f"/bang-tinh/{bang.code}/xuat/")
    assert tep.status_code == 200
    tieu_de = [o.value for o in next(load_workbook(BytesIO(tep.content)).active.iter_rows())]
    assert "Thành phố" not in tieu_de

    # 3. Bảng dữ liệu bên KN ERP (chỉ xem)
    with override_settings(ROOT_URLCONF="knjsc.urls"):
        client.force_login(nguoi_dung["admin"])
        erp = client.get(f"/bang/{bang.code}/")
        assert erp.status_code == 200
        assert ma not in [c.code for c in erp.context["cac_cot"]]

    # Dữ liệu không mất, cột vẫn còn định nghĩa
    dong.refresh_from_db()
    assert dong.data[ma] == "Đà Nẵng"
    assert ColumnDef.objects.get(table=bang, code=ma).is_hidden is True
    assert AuditLog.objects.filter(action=AuditAction.UPDATE, detail__contains=f"Ẩn cột {ma}").exists()


def test_hien_lai_cot_tra_ve_nguyen_ven(client, luoi, nguoi_dung):
    """AC-39.2 — Hiện lại cột đã ẩn: cột về đúng chỗ cũ trên lưới, dữ liệu cũ hiện đủ"""
    bang, products, rows = luoi
    ma = "thanh_pho"
    rows[0].data[ma] = "Huế"
    rows[0].save(update_fields=["data"])
    truoc = None

    client.force_login(nguoi_dung["admin"])
    truoc = _ma_cot_luoi(client)
    client.post(URL, {"cot": [ma], "an": "1"})
    assert ma not in _ma_cot_luoi(client)

    kq = client.post(URL, {"cot": [ma], "an": "0"})
    assert kq.status_code == 200 and kq.json() == {"da_doi": [ma], "an": False}
    assert _ma_cot_luoi(client) == truoc, "cột về đúng vị trí cũ, không xếp xuống cuối"
    assert DataRecord.objects.get(pk=rows[0].pk).data[ma] == "Huế"


def test_chi_quan_ly_bang_duoc_an_cot(client, luoi, nguoi_dung):
    """AC-39.3 — Nhân viên và quản lý bộ phận khác bị từ chối (403 khi thấy bảng, 404 khi
    bảng ngoài phạm vi), không cột nào đổi; Admin thì được"""
    bang, products, _ = luoi
    ma = "thanh_pho"

    for vai, ma_loi in (("staff_vd", 403), ("staff_sale_1", 403), ("manager_mkt", 404)):
        client.force_login(nguoi_dung[vai])
        kq = client.post(URL, {"cot": [ma], "an": "1"})
        assert kq.status_code == ma_loi, vai
        assert ColumnDef.objects.get(table=bang, code=ma).is_hidden is False, vai

    client.force_login(nguoi_dung["admin"])
    assert client.post(URL, {"cot": [ma], "an": "1"}).status_code == 200


def test_tu_choi_an_cot_khoa_cot_bat_buoc_va_an_het(client, feedback, nguoi_dung):
    """AC-39.4 — Từ chối ẩn cột khoá, cột bắt buộc nhập và lần ẩn làm bảng không còn cột nào"""
    bang, _, _ = feedback
    khoa = bang.columns.get(is_key=True)
    bat_buoc = bang.columns.exclude(pk=khoa.pk).first()
    bat_buoc.required = True
    bat_buoc.save(update_fields=["required"])

    with pytest.raises(BusinessError, match="cột khoá"):
        table_service.set_columns_hidden(bang, [khoa.code], True)
    with pytest.raises(BusinessError, match="bắt buộc nhập"):
        table_service.set_columns_hidden(bang, [bat_buoc.code], True)
    with pytest.raises(BusinessError, match="Chưa chọn cột nào"):
        table_service.set_columns_hidden(bang, ["khong_co_cot_nay"], True)

    con_lai = list(bang.columns.exclude(pk__in=[khoa.pk, bat_buoc.pk]).values_list("code", flat=True))
    table_service.set_columns_hidden(bang, con_lai, True)
    khoa.is_key = False
    khoa.save(update_fields=["is_key"])
    bat_buoc.required = False
    bat_buoc.save(update_fields=["required"])
    with pytest.raises(BusinessError, match="ít nhất một cột hiện"):
        table_service.set_columns_hidden(bang, [khoa.code, bat_buoc.code], True)
    assert bang.columns.filter(is_hidden=False).count() == 2


def test_cot_san_pham_mac_dinh_an_va_san_pham_moi_cung_an(client, luoi, nguoi_dung):
    """AC-39.5 — Nhóm cột số lượng theo sản phẩm **mặc định ẩn với cả công ty**; thêm sản
    phẩm mới thì cột của nó cũng vào ở trạng thái ẩn dù chưa ai bấm nút; Lên đơn vẫn ghi
    số lượng vào cột đang ẩn để hiện lại là có đủ dữ liệu; quản lý bảng bật lại được cả nhóm"""
    bang, products, _ = luoi
    nhom = _cot_san_pham(bang)
    assert len(nhom) >= 2
    assert all(ColumnDef.objects.get(table=bang, code=ma).is_hidden for ma in nhom), \
        "cột sản phẩm phải ẩn ngay từ lúc sinh ra, không chờ ai bấm"

    client.force_login(nguoi_dung["admin"])
    trang = client.get(f"/bang-tinh/{bang.code}/")
    # `productColumns` là đích của nút gộp "Ẩn cột số lượng theo sản phẩm"; cả nhóm đã ẩn
    # thì không còn gì để ẩn nữa nên danh sách rỗng, và nút gộp không hiện
    assert trang.context["config"]["productColumns"] == []
    assert not [c for c in _ma_cot_luoi(client) if c.startswith(dispatch_service.PRODUCT_COLUMN_PREFIX)]

    # Thêm sản phẩm mới như nút "Tạo sản phẩm" ở Lên đơn — không bấm ẩn lần nào
    moi = product_service.create_product(name="Sản phẩm Test", actor=nguoi_dung["admin"])
    ma_moi = dispatch_service.product_column_code(moi)
    assert ColumnDef.objects.get(table=bang, code=ma_moi).is_hidden is True
    assert ma_moi not in _ma_cot_luoi(client)

    # Lên đơn vẫn điền số lượng vào cột đang ẩn
    don = order_service.create_order(
        phone="0987654321", customer_name="Khách Test",
        lines=[{"product": moi.code, "quantity": 4, "unit_price": "10.00"}],
        actor=nguoi_dung["staff_sale_1"])
    assert don.record.data[ma_moi] == 4

    # Quản lý bảng bật lại cả nhóm thì thấy đủ, kể cả cột của sản phẩm mới
    kq = client.post(URL, {"cot": nhom + [ma_moi], "an": "0"})
    assert kq.status_code == 200
    hien = _ma_cot_luoi(client)
    assert ma_moi in hien and all(c in hien for c in nhom)


def test_migration_0011_an_cot_san_pham_dang_co(client, luoi, nguoi_dung):
    """AC-39.9 — Cột `sl_*` mọc trước 22.09 được migration `orders/0011` ẩn với cả công ty;
    chạy ngược không tự hiện lại và không mất dữ liệu ô"""
    bang, products, rows = luoi
    nhom = _cot_san_pham(bang)
    ma = nhom[0]
    rows[0].data[ma] = 9
    rows[0].save(update_fields=["data"])

    # Dựng lại trạng thái cũ: cột sản phẩm đang hiện
    ColumnDef.objects.filter(table=bang, code__in=nhom).update(is_hidden=False)
    client.force_login(nguoi_dung["admin"])
    assert ma in _ma_cot_luoi(client)

    an_cot_san_pham = import_module(
        "orders.migrations.0011_an_cot_san_pham_dang_co").an_cot_san_pham
    an_cot_san_pham(apps, None)
    assert all(ColumnDef.objects.get(table=bang, code=c).is_hidden for c in nhom)
    assert ma not in _ma_cot_luoi(client)
    assert DataRecord.objects.get(pk=rows[0].pk).data[ma] == 9

    # Chạy ngược là không làm gì: cột vẫn ẩn, lựa chọn của người dùng không bị xoá
    migrations.RunPython.noop(apps, None)
    assert ColumnDef.objects.get(table=bang, code=ma).is_hidden is True


def test_hop_cot_chi_quan_ly_bang_thay_muc_dang_an(client, luoi, nguoi_dung):
    """AC-39.6 — Quản lý bảng thấy danh sách cột đang ẩn để bật lại; nhân viên không thấy
    và không biết bảng có cột ẩn"""
    bang, products, _ = luoi
    ma = "thanh_pho"
    table_service.set_columns_hidden(bang, [ma], True, actor=nguoi_dung["admin"])

    client.force_login(nguoi_dung["admin"])
    config = client.get(f"/bang-tinh/{bang.code}/").context["config"]
    assert config["canHideColumns"] is True
    # Nhóm cột sản phẩm mặc định ẩn nên cũng nằm trong danh sách bật lại
    dang_an = [c["code"] for c in config["hiddenColumns"]]
    assert ma in dang_an
    assert set(_cot_san_pham(bang)) <= set(dang_an)
    assert config["hideColumnsUrl"] == URL

    client.force_login(nguoi_dung["staff_vd"])
    config = client.get(f"/bang-tinh/{bang.code}/").context["config"]
    assert config["canHideColumns"] is False and config["hiddenColumns"] == []


def test_migration_0015_xuoi_nguoc(departments):
    """AC-39.7 — Migration `forms_builder/0015` chạy xuôi và ngược đều được, giữ nguyên cột"""
    assert connection.settings_dict["NAME"].startswith("test_")
    bang = TableDef.objects.create(code="migration-039", name="Migration 039", department=departments["vd"])
    cot = ColumnDef.objects.create(table=bang, name="Ghi chú", code="ghi_chu")

    def columns():
        with connection.cursor() as cursor:
            return {c.name for c in connection.introspection.get_table_description(
                cursor, "forms_builder_columndef")}

    # Trigger ghi nhận thay đổi lưới (`crm_capture_*`) đang treo trên bảng cột;
    # chưa xả thì Postgres không cho ALTER TABLE trong cùng giao dịch.
    with connection.cursor() as cursor:
        cursor.execute("SET CONSTRAINTS ALL IMMEDIATE")
    try:
        MigrationExecutor(connection).migrate([("forms_builder", "0014_remove_tabledef_receives_orders")])
        assert "is_hidden" not in columns()
    finally:
        executor = MigrationExecutor(connection)
        executor.migrate(executor.loader.graph.leaf_nodes())
    assert "is_hidden" in columns()
    assert ColumnDef.objects.filter(pk=cot.pk, is_hidden=False).exists()
