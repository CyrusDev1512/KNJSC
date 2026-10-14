"""Báo cáo tổng hợp bản cũ không cộng lẫn loại tiền — TL-92 · P0-4.

Bảng chưa có nguồn báo cáo (`ReportSource`) thì màn hình rơi về bản tổng quát cũ (`summary_service`):
`configure_erp_reports` hỏng cả lệnh và hoàn lại, bảng nhập tay có cột Loại tiền, hoặc bảng chưa cấu hình
đứng đầu danh sách. Bản cũ trước đây cộng thẳng USD với CAD ở dòng tổng, ô số đầu trang, phụ chú kỳ trước,
Excel và thẻ Tổng quan — sai quy tắc bắt buộc 6 và ADR-046. Nay tách theo loại tiền như bản mới.
"""
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from forms_builder.meaning import FieldType
from forms_builder.models import ColumnDef
from forms_builder.services import record_service
from reports import aggregations
from reports.services import summary_service
from reports.tests.test_aggregations import bang_mkt  # noqa: F401 — fixture
from reports.tests.test_summary_scope import bang_bc_sale  # noqa: F401 — fixture

pytestmark = pytest.mark.django_db

THANG_8 = {"date_from": date(2026, 8, 1), "date_to": date(2026, 8, 31), "product": ""}

#: (người tạo, ngày, người bán, số đơn, doanh số, loại tiền)
DONG = [
    ("staff_sale_1", "2026-08-01", "Dũng", "2", "1234", "CAD"),
    ("staff_sale_1", "2026-08-01", "Dũng", "1", "4321", "USD"),
    ("staff_sale_2", "2026-08-02", "Hằng", "3", "1000", "CAD"),
    ("staff_sale_1", "2026-07-15", "Dũng", "1", "500", "CAD"),     # kỳ trước
]


def _cot_loai_tien(bang):
    ColumnDef.objects.create(table=bang, name="Loại tiền", code="loai_tien", field_type=FieldType.CHOICE,
                             options=["USD", "CAD"], order=6)


@pytest.fixture
def hai_loai_tien(bang_bc_sale, nguoi_dung):  # noqa: F811
    """Bảng Sale có cột Loại tiền, KHÔNG có nguồn báo cáo — đi đường cũ."""
    _cot_loai_tien(bang_bc_sale)
    for nguoi, ngay, ban, don, ds, tien in DONG:
        record_service.create_record(bang_bc_sale, {
            "ngay": ngay, "nguoi_ban": ban, "san_pham": "SP A", "so_don": don, "doanh_so": ds,
            "loai_tien": tien}, actor=nguoi_dung[nguoi])
    return bang_bc_sale


@pytest.mark.parametrize("tab", ["tong-hop", "nhan-vien", "san-pham"])
def test_duong_cu_tach_tong_theo_loai_tien_khong_cong_lan(hai_loai_tien, nguoi_dung, tab):
    """AC-46.1 — Bản cũ (bảng chưa có nguồn báo cáo) tách TỔNG CỘNG theo loại tiền, mỗi dòng nhóm một loại
    tiền, tổng chung không giữ số tiền cộng lẫn (TL-92)"""
    ctx = summary_service.build_context(nguoi_dung["manager_sale"], hai_loai_tien, tab=tab, **THANG_8)
    kq = ctx["kq"]

    assert kq.currency_key == "loai_tien"
    tong = dict(aggregations.total_rows(kq))
    assert list(tong) == ["USD", "CAD"]
    doanh_so = [i for i, c in enumerate(kq.columns) if c.code == "doanh_so"][0]
    assert (tong["USD"][doanh_so], tong["CAD"][doanh_so]) == (Decimal("4321"), Decimal("2234"))
    assert kq.totals["c_doanh_so"] is None and kq.totals["so_dong"] == 3
    assert all("loai_tien" in nhom for nhom in ctx["cac_nhom"])


def test_duong_cu_o_so_va_so_ky_truoc_khong_cong_lan(hai_loai_tien, nguoi_dung):
    """AC-46.1 — Ô số đầu trang của bản cũ ghi tiền theo từng loại ("4.321 USD · 2.234 CAD"), không so kỳ
    trước qua hai loại tiền; cột đếm vẫn so (TL-92)"""
    ctx = summary_service.build_context(nguoi_dung["manager_sale"], hai_loai_tien, tab="tong-hop",
                                        with_compare=True, **THANG_8)
    o = {x["nhan"]: x for x in ctx["o_so"]}
    chu = " ".join(x["gia_tri"] + " " + x.get("phu", {}).get("text", "") for x in ctx["o_so"])

    assert "6.555" not in chu and "6.055" not in chu and "5.555" not in chu
    assert o["Doanh số"]["gia_tri"] == "4.321 USD · 2.234 CAD"
    assert o["Doanh số"]["phu"]["chieu"] == "" and o["Doanh số"]["phu"]["khong_so"]
    assert o["Số đơn"]["phu"] == {"chieu": "len", "text": "+5"}


def test_duong_cu_xuat_excel_tong_theo_loai_tien(client, hai_loai_tien, nguoi_dung):
    """AC-46.4 — Excel của bản cũ có cột Loại tiền, TỔNG CỘNG mỗi loại tiền một dòng, số thô đúng như nhập,
    không ô nào cộng lẫn; phụ đề nói rõ không quy đổi (TL-92)"""
    client.force_login(nguoi_dung["manager_sale"])
    r = client.get("/bao-cao/tong-hop/xuat/?nguon=bc_sale&nhom=nhan-vien&tu=2026-08-01&den=2026-08-31")
    ws = load_workbook(BytesIO(r.content), data_only=True).active
    hang = list(ws.iter_rows(values_only=True))

    assert "không quy đổi" in (hang[1][0] or "")
    tieu_de = list(hang[3])
    assert "Loại tiền" in tieu_de
    tong = [h for h in hang[4:] if str(h[0] or "").startswith("Tổng cộng")]
    assert [h[0] for h in tong] == ["Tổng cộng · USD", "Tổng cộng · CAD"]
    cot_ds = tieu_de.index("Doanh số")
    assert [Decimal(str(h[cot_ds])) for h in tong] == [Decimal("4321"), Decimal("2234")]
    assert all(Decimal("6555") != Decimal(str(o)) for h in hang[4:] for o in h if isinstance(o, (int, float, Decimal)))
    cot_tien = tieu_de.index("Loại tiền")
    assert {h[cot_tien] for h in hang[4:] if h not in tong} == {"USD", "CAD"}


def test_trang_mac_dinh_khong_nguon_khong_cong_lan_loai_tien(client, hai_loai_tien, nguoi_dung):
    """AC-46.1 — Trang Báo cáo tổng hợp không ghi nguồn mà rơi vào bảng chưa cấu hình thì vẫn tách loại tiền,
    đúng phạm vi: Manager thấy cả bộ phận, Leader chỉ thấy team mình, bộ phận khác 403 (TL-92)"""
    client.force_login(nguoi_dung["manager_sale"])
    r = client.get("/bao-cao/tong-hop/?tu=2026-08-01&den=2026-08-31")
    html = r.content.decode()
    assert r.context["bang"].code == "bc_sale"
    assert "<th>Loại tiền</th>" in html and "Tổng cộng · USD" in html and "Tổng cộng · CAD" in html
    assert "6.555" not in html and "không quy đổi" in html

    client.force_login(nguoi_dung["leader_sale_1"])
    html = client.get("/bao-cao/tong-hop/?tu=2026-08-01&den=2026-08-31").content.decode()
    assert "Tổng cộng · CAD" in html and "1.234" in html
    assert "2.234" not in html and "Hằng" not in html

    client.force_login(nguoi_dung["staff_mkt"])
    assert client.get("/bao-cao/tong-hop/?nguon=bc_sale").status_code == 403


def test_duong_cu_dong_chua_co_loai_tien_vao_nhom_chua_ro(client, hai_loai_tien, nguoi_dung):
    """AC-46.2 — Bản cũ: dòng chưa có loại tiền thành nhóm "Chưa rõ", không cộng vào loại tiền nào, có cảnh
    báo nêu số dòng (TL-92)"""
    record_service.create_record(hai_loai_tien, {
        "ngay": "2026-08-03", "nguoi_ban": "Dũng", "san_pham": "SP A", "so_don": "1", "doanh_so": "99"},
        actor=nguoi_dung["staff_sale_1"])
    client.force_login(nguoi_dung["manager_sale"])
    html = client.get("/bao-cao/tong-hop/?nguon=bc_sale&tu=2026-08-01&den=2026-08-31").content.decode()

    assert "Tổng cộng · Chưa rõ" in html and "1 dòng chưa có loại tiền" in html
    assert "Tổng cộng · CAD" in html and "2.333" not in html


def test_duong_cu_tach_loai_tien_khong_them_truy_van(client, hai_loai_tien, nguoi_dung,
                                                    django_assert_max_num_queries):
    """AC-46.9 — Tách loại tiền ở bản cũ không thêm truy vấn ở đường trong bộ nhớ: Leader xem trang mặc định
    vẫn ≤ 10 lệnh (bài khoá, TL-92)"""
    client.force_login(nguoi_dung["leader_sale_1"])
    client.get("/bao-cao/tong-hop/")                     # lượt đầu ghi mốc phiên
    with django_assert_max_num_queries(10):
        assert client.get("/bao-cao/tong-hop/").status_code == 200


def test_the_tong_quan_kieu_cu_moi_loai_tien_mot_cot(client, bang_mkt, nguoi_dung):  # noqa: F811
    """AC-46.7 — Thẻ Thống kê Marketing của Tổng quan khi nguồn chưa cấu hình: mỗi chỉ tiêu một hàng, mỗi
    loại tiền một cột có tiêu đề mã tiền, không cộng hai loại tiền (TL-92)"""
    _cot_loai_tien(bang_mkt)
    for tien, cpqc, ds, mess, don in (("CAD", "7000", "1234", "10", "2"), ("USD", "1111", "4321", "5", "1")):
        record_service.create_record(bang_mkt, {
            "ngay": "2026-08-01", "marketer": "Minh", "san_pham": "SP A", "so_mess": mess, "cpqc": cpqc,
            "so_don": don, "doanh_so": ds, "loai_tien": tien}, actor=nguoi_dung["staff_mkt"])
    client.force_login(nguoi_dung["manager_mkt"])
    r = client.get("/?mkt_nguon=bc_mkt_test&mkt_tu=2026-08-01&mkt_den=2026-08-31")
    m = r.context["marketing"]["data"]
    html = r.content.decode()

    assert r.context["activity"]["configured"] is False
    assert m["state"] == "ready" and m["currencies"] == ["USD", "CAD"]
    chi_tieu = dict(m["metrics"])
    assert chi_tieu["CPQC"] == ["1.111", "7.000"] and chi_tieu["Doanh số"] == ["4.321", "1.234"]
    assert "8.111" not in html and "5.555" not in html
    assert '<th scope="col">USD</th>' in html
