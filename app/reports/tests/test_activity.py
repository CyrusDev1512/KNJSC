from core.identity import employee_code, identity_label
from datetime import date
from decimal import Decimal
import pytest
from forms_builder.models import DataRecord, TableDef
from reports.services import activity_service
from reports.models import ReportSource
from reports.tests.test_aggregations import bang_mkt, dong_mau
from reports.tests.test_mkt_excel import marketing_scope
from reports.tests.test_mkt_derived_revenue import mkt_source, van_don

pytestmark=pytest.mark.django_db


def test_sale_ratio_and_account_identity(bang_mkt,dong_mau,nguoi_dung):
    source=ReportSource.objects.create(table=bang_mkt,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    result=activity_service.build(nguoi_dung["manager_mkt"],source,group="person",start=date(2026,8,1),end=date(2026,8,31))
    assert result.totals["so_dong"]==4
    assert result.totals["conversion"] == Decimal(24)/Decimal(270)*100   # tỉ lệ hiện % (ADR-040)
    assert len(list(result.rows))==1  # Account ID, not different handwritten names.


def test_market_missing_is_preserved(bang_mkt,dong_mau,nguoi_dung):
    source=ReportSource.objects.create(table=bang_mkt,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    result=activity_service.build(nguoi_dung["manager_mkt"],source,group="market",start=date(2026,8,1),end=date(2026,8,31))
    assert list(result.rows)[0]["nhom"]=="Chưa xác định"


@pytest.mark.parametrize("role,count", [("staff_sale_1",1),("leader_sale_1",3),("manager_sale",4),("admin",4)])
def test_sale_scope_and_dashboard(client, marketing_scope, nguoi_dung, role, count):
    from io import BytesIO
    from openpyxl import load_workbook
    from reports import aggregations
    source=ReportSource.objects.create(table=marketing_scope.table,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    client.force_login(nguoi_dung[role])
    query={"nguon":source.table.code,"tu":"2026-08-01","den":"2026-08-31"}
    r=client.get("/bao-cao/tong-hop/",query)
    assert r.status_code==200
    result=r.context["result"]
    assert result.totals["so_dong"]==count
    exported=client.get("/bao-cao/tong-hop/xuat/",query)
    rows=list(load_workbook(BytesIO(exported.content),data_only=True).active.values)
    tong=aggregations.total_values(result)   # dòng TỔNG CỘNG của khối toàn kỳ (ADR-040): nhãn, ô trống, rồi số liệu
    dong_tong=next(r for r in rows if r[0] and str(r[0]).startswith("TỔNG CỘNG"))
    assert tuple(Decimal(str(v)) if v is not None else None for v in dong_tong[-len(tong):])==tuple(tong)
    d=client.get("/",{"sale_nguon":source.table.code,"tu":query["tu"],"den":query["den"]})
    block=next(b for b in d.context["activity"]["blocks"] if b["kind"]=="sale")
    assert block["ok"] and block["data"]["count"]==count


def test_foreign_activity_source(client, marketing_scope, nguoi_dung):
    ReportSource.objects.create(table=marketing_scope.table,kind="sale")
    client.force_login(nguoi_dung["staff_mkt"])
    for path in ("/bao-cao/tong-hop/", "/bao-cao/tong-hop/xuat/"):
        assert client.get(path,{"nguon":marketing_scope.table.code}).status_code==403


def test_zero_and_missing_inputs(bang_mkt, nguoi_dung):
    source=ReportSource.objects.create(table=bang_mkt,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    user=nguoi_dung["staff_mkt"]
    DataRecord.objects.create(table=bang_mkt,created_by=user,department=bang_mkt.department,val_date=date(2026,8,1),data={"so_mess":0,"so_don":2})
    result=activity_service.build(user,source)
    assert result.totals["conversion"] is None
    assert result.totals["__missing_revenue"] is None


@pytest.fixture
def delivery_source(marketing_scope, nguoi_dung):
    from orders.models import WaybillAssignment, WaybillItem, Product
    source=ReportSource.objects.create(table=marketing_scope.table,kind="delivery",columns={"market":"quoc_gia"})
    product=Product.objects.create(code="a",name="A")
    other=Product.objects.create(code="b",name="B")
    records=list(DataRecord.objects.filter(table=source.table).order_by("pk"))
    for i, row in enumerate(records):
        # Người tạo khác người được giao; hai sản phẩm không nhân số đơn.
        owner=nguoi_dung["staff_sale_1"] if i < 2 else nguoi_dung["staff_sale_2"]
        if i < 3: WaybillAssignment.objects.create(record=row,delivery=owner)
        row.data={"quoc_gia":"Canada" if i < 2 else "Hoa Kỳ", "trang_thai_vc":"Đang giao" if i < 2 else "Đã giao"}
        row.save(update_fields=["data"])
        WaybillItem.objects.create(record=row,product=product,quantity=2,unit_price=999)
        WaybillItem.objects.create(record=row,product=other,quantity=3,unit_price=999)
    return source


@pytest.mark.parametrize("role,count",[("staff_sale_1",2),("leader_sale_1",2),("manager_sale",4),("admin",4)])
def test_delivery_distinct_and_assigned_scope(delivery_source,nguoi_dung,role,count):
    source=delivery_source
    result=activity_service.build(nguoi_dung[role],source,group="person")
    assert result.totals["c_orders"]==count
    assert result.totals["c_quantity"]==count*5
    assert [c.label for c in result.columns]==["Số đơn","Số lượng sản phẩm"]
    rows=list(result.rows)
    if count==4: assert any(r["nhom"]=="Chưa phân công" for r in rows)


def test_delivery_filters_match_status_and_export(delivery_source,nguoi_dung,client):
    from io import BytesIO
    from openpyxl import load_workbook
    source=delivery_source
    user=nguoi_dung["admin"]
    result=activity_service.build(user,source,group="product",product="a",market="Canada")
    assert result.totals["c_orders"]==2 and result.totals["c_quantity"]==4
    assert len(list(result.rows))==1
    assert activity_service.shipping_status(user,source,product="a",market="Canada")==[{"label":"Đang giao","count":2}]
    client.force_login(user)
    r=client.get("/bao-cao/tong-hop/xuat/",{"nguon":source.table.code,"sp":"a","thi_truong":"Canada","tu":"2026-08-01","den":"2026-08-31"})
    assert r.status_code==200
    book=load_workbook(BytesIO(r.content),data_only=True)
    dong_tong=next(r for r in book.active.values if r[0] and str(r[0]).startswith("TỔNG CỘNG"))   # khối toàn kỳ (ADR-040)
    assert dong_tong[-2:]==(2,4)
    assert list(book["Trạng thái giao hàng"].values)[-1]==("Đang giao",2)


def test_configure_metadata_idempotent(departments):
    from django.core.management import call_command
    from forms_builder.models import FormDef, FormTableLink
    call_command("configure_erp_reports")
    call_command("configure_erp_reports")
    form=FormDef.objects.get(code="bc_sale_ngay")
    assert ReportSource.objects.filter(table=form.table).count()==1
    links=FormTableLink.objects.filter(form_field__form=form,column__code="thi_truong")
    assert links.count()==1 and links.get().form_field.required
    assert not DataRecord.objects.exists()


@pytest.mark.parametrize("market,day,valid",[("Canada","2026-08-01",True),("","2026-08-01",False),("Atlantis","2026-08-01",False),("Canada","2026-08-02",False)])
def test_submit_market_and_period(departments,nguoi_dung,market,day,valid):
    from django.core.management import call_command
    from forms_builder.models import FormDef
    from reports.services import daily_service
    from core.exceptions import BusinessError
    from orders.models import Product
    Product.objects.create(code="p",name="Sản phẩm P")
    call_command("configure_erp_reports")
    form=FormDef.objects.get(code="bc_sale_ngay")
    fields=list(form.ordered_fields())
    raw={"ngay":day,"san_pham":"Sản phẩm P","thi_truong":market,"so_mess":"10","so_don":"2","doanh_so":"100"}
    values={f.field.code:raw.get(f.link.column.code,"") for f in fields}
    if valid:
        report=daily_service.submit(form,values,report_date=date(2026,8,1),actor=nguoi_dung["staff_sale_1"],fields=fields)
        assert report.record.data["thi_truong"]==market
    else:
        with pytest.raises(BusinessError):
            daily_service.submit(form,values,report_date=date(2026,8,1),actor=nguoi_dung["staff_sale_1"],fields=fields)
        assert DataRecord.objects.count()==0


@pytest.mark.parametrize("path", ["/bao-cao/tong-hop/", "/bao-cao/lich-su/"])
def test_query_budget(client,marketing_scope,nguoi_dung,django_assert_max_num_queries,path):
    source=ReportSource.objects.create(table=marketing_scope.table,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    client.force_login(nguoi_dung["leader_sale_1"])
    params={"nguon":source.table.code,"tu":"2026-08-01","den":"2026-08-31"}
    client.get(path,params)
    with django_assert_max_num_queries(10):
        assert client.get(path,params).status_code==200



def test_marketing_exact_excel_formula(bang_mkt,dong_mau,nguoi_dung):
    from reports import aggregations
    source=ReportSource.objects.create(table=bang_mkt,kind="mkt",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","cost":"cpqc","market":"thi_truong"})
    result=activity_service.build(nguoi_dung["manager_mkt"],source,group="market")
    values=dict(zip([c.label for c in result.columns],aggregations.total_values(result)))
    assert values["CPQC"]==Decimal(500000)
    assert values["CPO"]==Decimal(500000)/24
    assert values["Giá Mess"]==Decimal(500000)/270
    # Hóa đơn ÷ DS Chốt (TT) theo nhãn (ADR-038, nhãn MKT theo ảnh ADR-040): không có vận đơn và Hóa đơn nên trống
    assert values["Hóa đơn/DS Chốt (TT)"] is None
    assert values["DS Chốt (TT)"] is None and values["Hóa đơn"] is None
    # Nguồn không ánh xạ Loại tiền thì cộng thô như cũ, không hậu tố ₫
    assert not result.converted and all(c.suffix != " ₫" for c in result.columns)


def test_dashboard_isolates_single_source_failure(client,marketing_scope,nguoi_dung,monkeypatch):
    from dashboard.services import dashboard_service
    ReportSource.objects.create(table=marketing_scope.table,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    client.force_login(nguoi_dung["admin"])
    def broken(*args,**kwargs): raise RuntimeError("injected test failure")
    monkeypatch.setattr(activity_service,"build",broken)
    response=client.get("/")
    assert response.status_code==200
    blocks={b["kind"]:b for b in response.context["activity"]["blocks"]}
    assert not blocks["sale"]["ok"]
    assert blocks["delivery"]["ok"] and blocks["delivery"]["data"]["state"]=="no_source"
    assert response.context["nhan_su"]["ok"]



def test_delivery_query_budget(client,delivery_source,nguoi_dung,django_assert_max_num_queries):
    client.force_login(nguoi_dung["leader_sale_1"])
    params={"nguon":delivery_source.table.code,"tu":"2026-08-01","den":"2026-08-31"}
    client.get("/bao-cao/tong-hop/",params)
    with django_assert_max_num_queries(10):
        assert client.get("/bao-cao/tong-hop/",params).status_code==200


@pytest.mark.parametrize("export",[False,True])
def test_legacy_link_uses_configured_report(client,marketing_scope,nguoi_dung,export):
    from urllib.parse import urlsplit,parse_qs
    ReportSource.objects.create(table=marketing_scope.table,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    client.force_login(nguoi_dung["manager_sale"])
    path="/bao-cao/hoat-dong/xuat/" if export else "/bao-cao/hoat-dong/"
    query={"nguon":marketing_scope.table.code,"nhom":"market","tu":"2026-08-01","den":"2026-08-31","sp":"SP1","thi_truong":"Canada","trang":"2"}
    response=client.get(path,query)
    assert response.status_code==302
    url=urlsplit(response.url)
    assert url.path==("/bao-cao/tong-hop/xuat/" if export else "/bao-cao/tong-hop/")
    assert parse_qs(url.query)=={k:[v] for k,v in query.items()}




def test_submission_does_not_fallback_from_foreign_form(client,marketing_scope,nguoi_dung):
    client.force_login(nguoi_dung["staff_mkt"])
    assert client.get("/bao-cao/",{"bieu_mau":marketing_scope.code}).status_code==403
    assert client.post("/bao-cao/",{"bieu_mau":marketing_scope.code,"ngay_bao_cao":"2026-08-01"}).status_code==403


def test_duplicate_names_remain_separate_accounts(marketing_scope,nguoi_dung):
    from org.models import UserProfile
    UserProfile.objects.filter(department=marketing_scope.department).update(full_name="Cùng tên")
    source=ReportSource.objects.create(table=marketing_scope.table,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    result=activity_service.build(nguoi_dung["admin"],source,group="person")
    rows=list(result.rows)
    assert len(rows)==4
    # Nhãn dòng là mã nhân sự, mỗi tài khoản một mã nên trùng họ tên vẫn tách dòng; họ tên không vào ô bảng
    assert len({row["nhom"] for row in rows})==4 and all("Cùng tên" not in row["nhom"] for row in rows)



def test_export_limit_before_workbook(client,bang_mkt,dong_mau,nguoi_dung,settings,monkeypatch):
    from reports import excel
    ReportSource.objects.create(table=bang_mkt,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    settings.EXPORT_MAX_ROWS=1
    client.force_login(nguoi_dung["manager_mkt"])
    def unexpected(*args,**kwargs): raise AssertionError("Must reject before allocating workbook")
    monkeypatch.setattr(excel,"build_workbook",unexpected)
    response=client.get("/bao-cao/tong-hop/xuat/",{"nguon":bang_mkt.code,"tu":"2026-08-01","den":"2026-08-31"})
    assert response.status_code==400
    assert "Thu hẹp bộ lọc" in response.content.decode()


@pytest.mark.parametrize("query",[{"nhom":"unknown"},{"tu":"2026-08-31","den":"2026-08-01"}])
def test_invalid_filters_show_error(client,marketing_scope,nguoi_dung,query):
    ReportSource.objects.create(table=marketing_scope.table,kind="sale",columns={"market":"thi_truong"})
    client.force_login(nguoi_dung["manager_sale"])
    response=client.get("/bao-cao/tong-hop/",{"nguon":marketing_scope.table.code,**query})
    assert response.status_code==400 and response.context["error"]


@pytest.mark.parametrize("old,new",[("tong-hop","day"),("nhan-vien","person"),("san-pham","product"),("thi-truong","market")])
def test_summary_preserves_old_group_links(client, marketing_scope, nguoi_dung, old, new):
    ReportSource.objects.create(table=marketing_scope.table, kind="sale", columns={"market":"thi_truong"})
    client.force_login(nguoi_dung["manager_sale"])
    response = client.get("/bao-cao/tong-hop/", {"nguon":marketing_scope.table.code,"nhom":old})
    assert response.status_code == 200
    assert response.context["params"]["group"] == new
    assert 'href="/bao-cao/hoat-dong/"' not in response.content.decode()


@pytest.mark.parametrize("role,persons", [
    ("staff_sale_1", {"staff_sale_1"}),
    ("leader_sale_1", {"staff_sale_1", "staff_sale_1b", "leader_sale_1"}),
    ("manager_sale", {"staff_sale_1", "staff_sale_1b", "staff_sale_2", "leader_sale_1"}),
    ("admin", {"staff_sale_1", "staff_sale_1b", "staff_sale_2", "leader_sale_1"}),
])
def test_day_view_shows_person_and_leader_in_scope(client, marketing_scope, nguoi_dung, role, persons):
    """AC-22.10 — Cách xem Tổng hợp nhóm theo ngày × nhân sự: cột Nhân sự và Leader ngay sau Ngày, Leader tra từ
    Team.leader; Staff chỉ thấy dòng của mình, Leader team mình, Manager cả bộ phận; lọc nhân sự thu hẹp bảng;
    Excel cùng cột; lọc nhân sự ngoài phạm vi bị 403"""
    from io import BytesIO
    from openpyxl import load_workbook
    source=ReportSource.objects.create(table=marketing_scope.table,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    client.force_login(nguoi_dung[role])
    query={"nguon":source.table.code,"tu":"2026-08-01","den":"2026-08-31"}
    r=client.get("/bao-cao/tong-hop/",query)
    assert r.status_code==200 and r.context["result"].show_person
    rows=[row for row in r.context["rows"] if row["kind"]=="row"]   # bỏ dòng Tổng ngày (AC-22.15)
    # Ngày × nhân sự: mỗi người một DÒNG riêng, ngày lặp lại (AC-22.14, chủ dự án 19.09)
    assert [row["nhom"] for row in rows]==["01.08.2026"]*len(persons)
    # Ô bảng chỉ mã nhân sự, không họ tên (ADR-037 bổ sung 18.09); ô chọn Nhân sự vẫn `MÃ · Họ tên`
    assert {row["person"] for row in rows}=={employee_code(nguoi_dung[name]) for name in persons}
    assert all(nguoi_dung[name].profile.full_name not in row["person"] for row in rows for name in persons)
    leaders={row["leader"] for row in rows}
    assert employee_code(nguoi_dung["leader_sale_1"]) in leaders    # Team.leader của Sale 1
    assert all(nguoi_dung["leader_sale_1"].profile.full_name not in row["leader"] for row in rows)
    assert (employee_code(nguoi_dung["leader_sale_2"]) in leaders) == ("staff_sale_2" in persons)
    assert {identity_label(nguoi_dung[name]) for name in persons} <= {p["label"] for p in r.context["people"]}
    assert r.context["label_span"]==4
    html=r.content.decode()
    assert 'class="report-identity id-nhan-su" data-pos="3">Nhân sự</th>' in html and 'id-leader report-identity-edge" data-pos="4">Leader</th>' in html   # cột định danh ghim (AC-22.13)
    # Lọc theo nhân sự: chỉ còn dòng của người đó; người ngoài phạm vi bị chặn
    me=nguoi_dung["staff_sale_1"].pk
    r2=client.get("/bao-cao/tong-hop/",{**query,"nhan_su":me})
    dong2=[row for row in r2.context["rows"] if row["kind"]=="row"]
    assert [row["person"] for row in dong2]==[employee_code(nguoi_dung["staff_sale_1"])] and dong2[0]["leader"]==employee_code(nguoi_dung["leader_sale_1"])
    outsider=nguoi_dung["staff_sale_2"].pk
    r3=client.get("/bao-cao/tong-hop/",{**query,"nhan_su":outsider})
    assert (r3.status_code==403) == ("staff_sale_2" not in persons)
    book=load_workbook(BytesIO(client.get("/bao-cao/tong-hop/xuat/",query).content),data_only=True)
    sheet=list(book.active.values)                                  # khối toàn kỳ theo nhân sự (ADR-040)
    assert sheet[4][:4]==("STT","Team","Nhân sự","Leader") and str(sheet[5][0]).startswith("TỔNG CỘNG")
    nguoi_excel=[d for d in sheet[6:] if d and d[0] is not None]
    assert len(nguoi_excel)==len(persons)                           # Excel cũng mỗi người một dòng
    assert {d[2] for d in nguoi_excel}=={employee_code(nguoi_dung[name]) for name in persons}
    assert employee_code(nguoi_dung["leader_sale_1"]) in {d[3] for d in nguoi_excel}
    assert nguoi_dung["staff_sale_1"].profile.full_name not in nguoi_excel[0][2]   # Excel cũng chỉ mã
    ngay=list(book["Theo ngay"].values)                              # mỗi ngày một khối: tiêu đề, cột, TỔNG CỘNG, người
    assert ngay[0][0]=="Ngày 01.08.2026" and ngay[1][:4]==("STT","Team","Nhân sự","Leader") and ngay[2][0]=="TỔNG CỘNG"


def test_day_view_groups_by_date_and_person(client, marketing_scope, nguoi_dung):
    """AC-22.14 — Cách xem Tổng hợp nhóm theo ngày × nhân sự: mỗi người một HÀNG riêng như ảnh mẫu,
    ngày lặp lại ở từng hàng; tổng trong bộ lọc không đổi; Excel cũng mỗi người một dòng"""
    from io import BytesIO
    from openpyxl import load_workbook
    source=ReportSource.objects.create(table=marketing_scope.table,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    client.force_login(nguoi_dung["admin"])
    query={"nguon":source.table.code,"tu":"2026-08-01","den":"2026-08-31"}
    r=client.get("/bao-cao/tong-hop/",query)
    rows=[row for row in r.context["rows"] if row["kind"]=="row"]
    # Bốn người cùng nộp ngày 01.08 → bốn hàng, không phải một hàng gộp
    assert len(rows)==4 and {row["nhom"] for row in rows}=={"01.08.2026"}
    assert len({row["person"] for row in rows})==4
    assert all(isinstance(row["person"],str) and "," not in row["person"] for row in rows)
    # Dòng Tổng vẫn tính trên toàn bộ kết quả, không đổi vì tách hàng
    assert r.context["result"].totals["so_dong"]==4
    # Mỗi hàng chỉ mang số của chính người đó
    assert [row["cells"][0] for row in rows]==["10"]*4
    sheet=list(load_workbook(BytesIO(client.get("/bao-cao/tong-hop/xuat/",query).content),data_only=True).active.values)
    nguoi_excel=[d for d in sheet[6:] if d and d[0] is not None]   # sau tiêu đề, hàng cột, TỔNG CỘNG của khối toàn kỳ
    assert len(nguoi_excel)==4 and len({d[2] for d in nguoi_excel})==4


def test_day_blocks_have_day_subtotal_and_stt(client, bang_mkt, mkt_source, van_don, nguoi_dung):
    """AC-22.15 — Cách xem Tổng hợp chia khối theo ngày như ảnh mẫu: mỗi ngày có dòng Tổng ngày đứng đầu,
    số của nó bằng tổng các dòng con và cột tính được tính lại từ tổng (không phải trung bình), Doanh thu
    suy ra cộng theo ngày; cột STT đếm lại từ 1 trong từng ngày; dòng Tổng trong bộ lọc không đổi; Excel cùng khối"""
    from io import BytesIO
    from decimal import Decimal
    from openpyxl import load_workbook
    from reports import aggregations
    from reports.tests.test_mkt_derived_revenue import _bao_cao
    A, B = van_don["A"], van_don["B"]
    _bao_cao(bang_mkt, A, "2026-08-01", "SP1", hoa_don="8")
    _bao_cao(bang_mkt, B, "2026-08-01", "SP2", hoa_don="2")
    _bao_cao(bang_mkt, A, "2026-08-02", "SP1", hoa_don="5")
    client.force_login(nguoi_dung["manager_mkt"])
    query={"nguon":bang_mkt.code,"tu":"2026-08-01","den":"2026-08-02"}
    r=client.get("/bao-cao/tong-hop/",query)
    assert r.status_code==200
    rows=r.context["rows"]
    nhan=[(row["kind"], row["nhom"], row.get("stt","")) for row in rows]
    # Hai khối: 02.08 (một người) rồi 01.08 (hai người), STT đếm lại từ 1 mỗi ngày
    assert nhan==[("subtotal","02.08.2026",""),("row","02.08.2026",1),
                  ("subtotal","01.08.2026",""),("row","01.08.2026",1),("row","01.08.2026",2)]
    cot=[c.label for c in r.context["result"].columns]
    def o(row, nhan_cot):
        return row["cells"][cot.index(nhan_cot)]
    khoi = rows[2]                                   # Tổng ngày 01.08
    con = [rows[3], rows[4]]
    # Số cộng được: Tổng ngày = tổng hai dòng con; tiền đã quy ₫ (CAD × 17.500, ADR-040)
    assert o(khoi,"Số Mess")=="20" and [o(d,"Số Mess") for d in con]==["10","10"]
    assert o(khoi,"Hóa đơn")=="175.000 ₫" and {o(d,"Hóa đơn") for d in con}=={"140.000 ₫","35.000 ₫"}
    # DS Chốt (TT) của ngày = tổng hai marketer (60+40 của A, 200 của B) = 300 CAD
    assert o(khoi,"DS Chốt (TT)")=="5.250.000 ₫"
    # Cột tính lại từ tổng, không phải trung bình các dòng con
    assert o(khoi,"Hóa đơn/DS Chốt (TT)")==aggregations.format_number(Decimal(175000)/Decimal(5250000), 4)
    # Dòng Tổng trong bộ lọc không đổi
    assert r.context["result"].totals["so_dong"]==3
    tong=dict(zip(cot, aggregations.total_cells(r.context["result"])))
    assert tong["Số Mess"]=="30" and tong["DS Chốt (TT)"]=="5.687.500 ₫"
    html=r.content.decode()
    # Mỗi ngày một bảng riêng (ADR-040): tiêu đề ngày trên bảng, TỔNG CỘNG ngay dưới hàng tiêu đề cột, STT ở cột đầu
    assert html.count('class="report-block report-block-day"')==2 and '<h3>01.08.2026</h3>' in html
    assert 'data-pos="1" colspan="4">TỔNG CỘNG</th>' in html and 'class="report-identity id-stt" data-pos="1">1</th>' in html
    assert 'class="report-block report-block-period"' in html and 'TỔNG CỘNG · toàn kỳ</th>' in html
    # Excel: sheet "Theo ngay" cùng khối — tiêu đề ngày, hàng tiêu đề cột, TỔNG CỘNG, dòng người có STT
    ngay=list(load_workbook(BytesIO(client.get("/bao-cao/tong-hop/xuat/",query).content),data_only=True)["Theo ngay"].values)
    assert [d[0] for d in ngay if d and d[0] is not None]==["Ngày 02.08.2026","STT","TỔNG CỘNG",1,
                                                              "Ngày 01.08.2026","STT","TỔNG CỘNG",1,2]


def test_metric_colours_against_filter_total(client, bang_mkt, mkt_source, van_don, nguoi_dung):
    """AC-22.16 — Tô màu chỉ tiêu: cột chỉ số quan trọng có nền riêng; ô tỉ lệ so với dòng Tổng trong bộ
    lọc theo chiều tốt của từng chỉ tiêu (CPO, Giá Mess thấp là đạt), lệch trong biên thì để trơn; cột
    cộng và chỉ tiêu chưa rõ chiều thì không tô; dòng Tổng là mốc nên không tô đạt/kém"""
    from reports import aggregations
    from reports.constants import FOCUS_METRICS, METRIC_DIRECTION
    from reports.tests.test_mkt_derived_revenue import _bao_cao
    A, B = van_don["A"], van_don["B"]
    # A: 20 mess, 8 đơn, CPQC 100 → CPO 12,5 · Giá Mess 5 ; B: 20 mess, 2 đơn, CPQC 100 → CPO 50 · Giá Mess 5
    # Tổng: 40 mess, 10 đơn, CPQC 200 → CPO 20 · Giá Mess 5
    _bao_cao(bang_mkt, A, "2026-08-01", "SP1", mess=20, don=8, cpqc="100")
    _bao_cao(bang_mkt, B, "2026-08-01", "SP2", mess=20, don=2, cpqc="100")
    client.force_login(nguoi_dung["manager_mkt"])
    query={"nguon":bang_mkt.code,"tu":"2026-08-01","den":"2026-08-01"}
    r=client.get("/bao-cao/tong-hop/",query)
    cot=[c.label for c in r.context["result"].columns]
    dong={row["person"]: row for row in r.context["rows"] if row["kind"]=="row"}
    def lop(row, nhan):
        return row["cells"][cot.index(nhan)].lop
    ma_a, ma_b = employee_code(A), employee_code(B)
    # CPO càng THẤP càng tốt: A 12,5 dưới mốc 20 là đạt; B 50 là cảnh báo. Cả hai có nền cột chỉ số.
    assert "o-chi-so" in lop(dong[ma_a],"CPO") and "o-tot" in lop(dong[ma_a],"CPO")
    assert "o-canh-bao" in lop(dong[ma_b],"CPO") and "o-tot" not in lop(dong[ma_b],"CPO")
    # Giá Mess của cả hai bằng đúng mốc → trong biên, chỉ có nền cột, không màu đạt/kém
    assert lop(dong[ma_a],"Giá Mess")=="o-chi-so" and lop(dong[ma_b],"Giá Mess")=="o-chi-so"
    # Cột cộng không tô: mốc là tổng mọi dòng nên dòng nào cũng nhỏ hơn
    assert "Số đơn" not in METRIC_DIRECTION and "Số Mess" not in METRIC_DIRECTION
    assert lop(dong[ma_a],"Số đơn")=="" and lop(dong[ma_a],"Số Mess")==""
    # Chỉ tiêu chưa rõ chiều cũng không tô
    assert "invoice_revenue" not in METRIC_DIRECTION and lop(dong[ma_a],"Hóa đơn/DS Chốt (TT)")==""
    # Dòng Tổng là mốc: không màu đạt/kém, vẫn giữ nền cột chỉ số
    tong=aggregations.total_cells(r.context["result"])
    assert all("o-tot" not in c.lop and "o-canh-bao" not in c.lop for c in tong)
    assert "o-chi-so" in tong[cot.index("CPO")].lop
    # Ô vẫn là chuỗi hiển thị như cũ — mọi chỗ so sánh không đổi
    assert dong[ma_a]["cells"][cot.index("Số Mess")]=="20"
    html=r.content.decode()
    assert '<th scope="col" class="o-chi-so">CPO</th>' in html   # nền cột ở tiêu đề
    assert 'class="o-chi-so o-tot"' in html and 'class="o-chi-so o-canh-bao"' in html
    # Chỉ số quan trọng khoá theo mã (ADR-040); nhãn MKT theo ảnh mẫu, kể cả Tỉ lệ chốt và Tỉ lệ chốt (TT)
    nhan={"conversion":"Tỉ lệ chốt","conversion_tt":"Tỉ lệ chốt (TT)","cpo":"CPO","mess_cost":"Giá Mess","cost_sales":"CPQC/DS Chốt"}
    assert {nhan[m] for m in FOCUS_METRICS} <= set(cot)


def test_day_view_pages_by_hundred(client, marketing_scope, nguoi_dung):
    """AC-22.11 — Báo cáo hoạt động phân trang mặc định 100 nhóm mỗi trang; dòng tổng trong bộ lọc vẫn tính trên
    toàn bộ kết quả, không theo trang"""
    source=ReportSource.objects.create(table=marketing_scope.table,kind="sale",columns={"mess":"so_mess","orders":"so_don","sales":"doanh_so","market":"thi_truong"})
    client.force_login(nguoi_dung["admin"])
    r=client.get("/bao-cao/tong-hop/",{"nguon":source.table.code,"tu":"2026-08-01","den":"2026-08-31"})
    assert r.context["trang"].paginator.per_page==100 and r.context["moi_trang"]==100
    assert r.context["result"].totals["so_dong"]==4
    r25=client.get("/bao-cao/tong-hop/",{"nguon":source.table.code,"tu":"2026-08-01","den":"2026-08-31","moi_trang":25})
    assert r25.context["moi_trang"]==25
