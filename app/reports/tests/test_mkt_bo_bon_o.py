"""Form báo cáo Marketing bỏ bốn ô — ADR-048 (chủ dự án 03.10.2026).

"Bỏ cả bốn ô": Sản phẩm, Thị trường, Tệp khách hàng, Loại tiền rời form Nộp báo cáo Marketing; cột và dữ liệu cũ
giữ nguyên; tiền vẫn là VND (ADR-047). Màn báo cáo đi theo: nguồn MKT không còn ba bộ lọc theo sản phẩm, thị
trường, tệp; bảng toàn VND không còn cột Loại tiền chỉ lặp một chữ. Báo cáo Sale và Vận đơn không đổi.
"""
from datetime import date
from io import BytesIO

import pytest
from django.core.management import call_command
from openpyxl import load_workbook

from forms_builder.models import ColumnDef, DataRecord, FormDef
from forms_builder.services import record_service
from orders.models import Product
from reports import aggregations
from reports.management.commands.configure_erp_reports import configure_source
from reports.models import DailyReport, ReportSource
from reports.tests.test_aggregations import bang_mkt  # noqa: F401
from reports.tests.test_mkt_derived_revenue import _bao_cao, mkt_source  # noqa: F401

pytestmark = pytest.mark.django_db

BON_O = ("san_pham", "thi_truong", "tep_khach_hang", "loai_tien")
KY = {"tu": "2026-08-01", "den": "2026-08-03"}


def _cot_tren_form(form):
    return {f.link.column.code for f in form.ordered_fields() if getattr(f, "link", None)}


def _payload(form, **theo_cot):
    out = {"bieu_mau": form.code}
    for f in form.ordered_fields():
        cot = getattr(getattr(f, "link", None), "column", None)
        if cot is not None and cot.code in theo_cot:
            out[f.field.code] = theo_cot[cot.code]
    return out


def _o_trong_bang(xlsx):
    sheet = load_workbook(BytesIO(xlsx)).active
    return [str(o) for dong in sheet.iter_rows(values_only=True) for o in dong if o is not None]


def test_form_mkt_bo_bon_o_sale_giu_nguyen(client, bang_mkt, mkt_source, nguoi_dung):  # noqa: F811
    """AC-48.1 — Sau `configure_erp_reports`, form Nộp báo cáo Marketing không còn Sản phẩm, Thị trường, Tệp khách
    hàng, Loại tiền (chạy lại vẫn không đưa về); bốn cột và ánh xạ market, currency, segment của nguồn giữ nguyên
    cho dòng cũ; trang nộp không còn ô nào của bốn cột; form Sale vẫn bắt buộc Sản phẩm và Thị trường"""
    form = mkt_source.table.forms.get()
    assert not set(BON_O) & _cot_tren_form(form)
    configure_source(bang_mkt, "mkt")
    assert not set(BON_O) & _cot_tren_form(form)
    assert ColumnDef.objects.filter(table=bang_mkt, code__in=BON_O).count() == 4
    mapping = ReportSource.objects.get(table=bang_mkt).columns
    assert {"market", "currency", "segment"} <= set(mapping)
    client.force_login(nguoi_dung["staff_mkt"])
    html = client.get("/bao-cao/", {"bieu_mau": form.code}).content.decode()
    for ma in BON_O:
        assert f"{ma}\"" not in html, ma
    assert "report-currency-map" not in html and "data-report-currency" not in html
    # Sale không đổi
    call_command("configure_erp_reports")
    sale = FormDef.objects.get(code="bc_sale_ngay")
    bat_buoc = {f.link.column.code for f in sale.ordered_fields() if f.required and getattr(f, "link", None)}
    assert {"san_pham", "thi_truong"} <= bat_buoc


def test_nop_va_sua_mkt_khong_can_thi_truong(client, bang_mkt, mkt_source, nguoi_dung):  # noqa: F811
    """AC-48.2 — Nộp báo cáo Marketing chỉ với số (không Thị trường, Sản phẩm) thì lưu được, dòng mang Loại tiền
    VND; Sửa báo cáo cũng không đòi Thị trường; cờ bắt buộc cấp cột của bốn cột trên bảng MKT (nếu từng bật tay)
    bị gỡ khi chạy lệnh cấu hình; dòng tạo qua đường nhập bảng không có Thị trường cũng ra VND; Sale nộp thiếu
    Thị trường vẫn bị từ chối"""
    ColumnDef.objects.filter(table=bang_mkt, code__in=BON_O).update(required=True)
    configure_source(bang_mkt, "mkt")
    assert not ColumnDef.objects.filter(table=bang_mkt, code__in=BON_O, required=True).exists()
    form = mkt_source.table.forms.get()
    client.force_login(nguoi_dung["staff_mkt"])
    r = client.post("/bao-cao/", _payload(form, so_mess="10", cpqc="13250000", so_don="2", doanh_so="100"))
    assert r.status_code == 302, r.content.decode()[:400]
    bao_cao = DailyReport.objects.get()
    assert bao_cao.record.data["loai_tien"] == "VND"
    assert not bao_cao.record.data.get("thi_truong") and not bao_cao.record.data.get("san_pham")
    # Sửa báo cáo (quản lý) không đòi Thị trường
    client.force_login(nguoi_dung["manager_mkt"])
    trang = client.get(f"/bao-cao/{bao_cao.pk}/sua/")
    assert trang.status_code == 200
    cpqc = next(f.field.code for f in form.ordered_fields() if f.link.column.code == "cpqc")
    r = client.post(f"/bao-cao/{bao_cao.pk}/sua/", {cpqc: "12000000", "version": trang.context["version"]})
    assert r.status_code == 302, r.content.decode()[:400]
    bao_cao.record.refresh_from_db()
    assert str(bao_cao.record.data["cpqc"]).startswith("12000000") and bao_cao.record.data["loai_tien"] == "VND"
    # Đường nhập bảng (cùng hàm Nhập tệp gọi từng dòng): không Thị trường vẫn ra VND
    dong = record_service.create_record(bang_mkt, {"ngay": "2026-08-02", "so_mess": 5, "cpqc": "1", "so_don": 1,
                                                   "doanh_so": "10"},
                                        actor=nguoi_dung["staff_mkt"], system_day=date(2026, 8, 2))
    assert DataRecord.objects.get(pk=dong.pk).data["loai_tien"] == "VND"
    # Sale vẫn bắt chọn Thị trường
    call_command("configure_erp_reports")
    sale = FormDef.objects.get(code="bc_sale_ngay")
    client.force_login(nguoi_dung["manager_sale"])
    r = client.post("/bao-cao/", _payload(sale, so_mess="10", so_don="2", doanh_so="100", san_pham="SP1"))
    # Từ chối ngay ở bước suy loại tiền theo Thị trường (như trước ADR-048) — không lưu gì
    assert r.status_code == 200 and r.context["loi"] == ["Chọn quốc gia hợp lệ để xác định loại tiền."]
    assert DailyReport.objects.filter(form=sale).count() == 0


def test_bo_loc_mkt_an_va_tham_so_cu_bi_bo_qua(client, bang_mkt, mkt_source, nguoi_dung):  # noqa: F811
    """AC-48.3 — Nguồn Marketing: Báo cáo tổng hợp và Bảng dữ liệu không còn bộ lọc Sản phẩm, Thị trường, Tệp khách
    hàng; URL cũ có `sp`, `thi_truong`, `tep` (kể cả giá trị lạ) vẫn mở được (200, không lỗi 400), không có chip,
    không lọc mất dòng, Excel không ghi chúng ở dòng phụ; nguồn Sale vẫn đủ ba bộ lọc"""
    A = nguoi_dung["staff_mkt"]
    _bao_cao(bang_mkt, A, "2026-08-01", "SP1")
    _bao_cao(bang_mkt, A, "2026-08-02", "SP2")
    client.force_login(nguoi_dung["manager_mkt"])
    cu = {**KY, "nguon": bang_mkt.code, "sp": "SP9", "thi_truong": "Klingon", "tep": "Klingon"}
    r = client.get("/bao-cao/tong-hop/", cu)
    assert r.status_code == 200
    html = r.content.decode()
    for ten in ('name="sp"', 'name="thi_truong"', 'name="tep"'):
        assert ten not in html, ten
    assert not {c["label"] for c in r.context["chips"]} & {"Sản phẩm", "Thị trường", "Tệp"}
    assert aggregations.row_count(r.context["result"]) == 2
    xlsx = client.get("/bao-cao/tong-hop/xuat/", cu)
    assert xlsx.status_code == 200
    assert not any("Sản phẩm:" in o or "Tệp khách hàng:" in o for o in _o_trong_bang(xlsx.content))
    # Bảng dữ liệu dạng báo cáo và tệp xuất của nó
    r = client.get(f"/bang/{bang_mkt.code}/", {k: v for k, v in cu.items() if k != "nguon"})
    assert r.status_code == 200
    html = r.content.decode()
    for ten in ('name="sp"', 'name="thi_truong"', 'name="tep"'):
        assert ten not in html, ten
    xuat = client.get(f"/bang/{bang_mkt.code}/xuat/", {k: v for k, v in cu.items() if k != "nguon"})
    assert xuat.status_code == 200 and xuat["Content-Type"].startswith("application/vnd.openxmlformats")
    # Sale vẫn đủ bộ lọc
    ReportSource.objects.filter(pk=mkt_source.pk).update(kind="sale")
    client.force_login(nguoi_dung["admin"])
    html = client.get("/bao-cao/tong-hop/", {**KY, "nguon": bang_mkt.code}).content.decode()
    for ten in ('name="sp"', 'name="thi_truong"', 'name="tep"'):
        assert ten in html, ten


def test_bang_mkt_toan_vnd_khong_cot_loai_tien(client, bang_mkt, mkt_source, nguoi_dung):  # noqa: F811
    """AC-48.4 — Báo cáo Marketing toàn VND: bảng không có cột Loại tiền (ẩn theo ADR-047 bổ sung), cột đứng yên cuối
    là Nhân sự (bóng mép ở ô Nhân sự); đơn vị ghi một lần: câu "Mọi số tiền là tiền Việt (₫), không quy đổi" và chữ
    "Tiền: ₫" ở hàng tiêu đề; Excel không có cột Loại tiền nhưng giữ nhãn "TỔNG CỘNG · toàn kỳ · VND"; dữ liệu lỡ lẫn
    loại tiền khác (nhãn ghi tay) thì cột vẫn ẩn (chủ dự án 03.10.2026) nhưng hàng tiêu đề không còn ghi "Tiền: ₫" nữa"""
    A = nguoi_dung["staff_mkt"]
    _bao_cao(bang_mkt, A, "2026-08-01", "SP1")
    lech = _bao_cao(bang_mkt, A, "2026-08-02", "SP1")
    client.force_login(nguoi_dung["manager_mkt"])
    r = client.get("/bao-cao/tong-hop/", {**KY, "nguon": bang_mkt.code})
    html = r.content.decode()
    assert r.context["result"].fixed_currency == "VND"
    assert ">Loại tiền</th>" not in html and "id-tien" not in html
    assert 'class="report-identity id-nhan-su report-identity-edge"' in html
    assert "Mọi số tiền là tiền Việt (₫), không quy đổi" in html and "Tiền: ₫" in html
    o = _o_trong_bang(client.get("/bao-cao/tong-hop/xuat/", {**KY, "nguon": bang_mkt.code}).content)
    assert "Loại tiền" not in o and "TỔNG CỘNG · toàn kỳ · VND" in o
    # Một dòng mang nhãn khác (dữ liệu ghi tay, không qua form): cột vẫn ẩn như mọi báo cáo MKT, nhưng hàng tiêu đề
    # không còn khẳng định mọi số là ₫
    DataRecord.all_objects.filter(pk=lech.pk).update(data={**lech.data, "loai_tien": "CAD"})
    r = client.get("/bao-cao/tong-hop/", {**KY, "nguon": bang_mkt.code})
    assert r.context["result"].fixed_currency == ""
    html = r.content.decode()
    assert ">Loại tiền</th>" not in html and "id-tien" not in html and "Tiền: ₫" not in html


def test_form_sale_bo_ngay_ra_don(client, nguoi_dung):
    """AC-48.7 — Sau `configure_erp_reports`, form Nộp báo cáo Sale không còn ô Ngày ra đơn (chạy lại vẫn không đưa về);
    cột Ngày ra đơn và dữ liệu cũ của cột giữ nguyên; cờ bắt buộc cấp cột (nếu từng bật tay) bị gỡ; trang nộp Sale không
    còn ô đó, nộp Sale đủ các ô còn lại vẫn lưu được"""
    Product.objects.get_or_create(code="sp1", defaults={"name": "SP1"})
    call_command("configure_erp_reports")
    sale = FormDef.objects.get(code="bc_sale_ngay")
    nguoi_ban = nguoi_dung["staff_sale_1"]
    cu = record_service.create_record(sale.table, {"ngay": "2026-08-01", "san_pham": "SP1", "so_mess": 5, "so_don": 1,
                                                   "doanh_so": "10", "ngay_ra_don": "2026-07-30", "thi_truong": "Canada"},
                                      actor=nguoi_ban, system_day=date(2026, 8, 1))
    assert "ngay_ra_don" not in _cot_tren_form(sale)
    cot = ColumnDef.objects.get(table=sale.table, code="ngay_ra_don")
    ColumnDef.objects.filter(pk=cot.pk).update(required=True)
    call_command("configure_erp_reports")
    assert "ngay_ra_don" not in _cot_tren_form(FormDef.objects.get(code="bc_sale_ngay"))
    assert not ColumnDef.objects.get(pk=cot.pk).required
    assert DataRecord.objects.get(pk=cu.pk).data["ngay_ra_don"] == "2026-07-30"     # dữ liệu cũ giữ
    client.force_login(nguoi_ban)
    html = client.get("/bao-cao/", {"bieu_mau": sale.code}).content.decode()
    assert 'ngay_ra_don"' not in html and "Ngày ra đơn" not in html
    r = client.post("/bao-cao/", _payload(sale, so_mess="10", so_don="2", doanh_so="100", san_pham="SP1",
                                          thi_truong="Canada"))
    assert r.status_code == 302, r.content.decode()[:400]
    assert DailyReport.objects.filter(form=sale).count() == 1
