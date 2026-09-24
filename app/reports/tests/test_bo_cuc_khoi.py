"""ADR-042 đợt 2 — bố cục khối như ảnh mẫu: khối toàn kỳ theo nhân sự, mỗi ngày một bảng, Gộp."""
from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from core.identity import employee_code
from reports import aggregations
from reports.tests.test_aggregations import bang_mkt  # noqa: F401
from reports.tests.test_mkt_derived_revenue import _bao_cao, mkt_source, van_don  # noqa: F401

pytestmark = pytest.mark.django_db


def _o(row, cot, nhan):
    return row["cells"][cot.index(nhan)]


def test_bang_toan_ky_va_moi_ngay_mot_bang(client, bang_mkt, mkt_source, van_don, nguoi_dung):
    """AC-42.6 — Cách xem Tổng hợp: khối toàn kỳ theo nhân sự đứng đầu (mỗi người một dòng cộng cả
    kỳ, STT, Team, Leader, TỔNG CỘNG toàn kỳ ngay dưới tiêu đề cột), rồi mỗi ngày một bảng riêng mới
    nhất trước không có cột Ngày, TỔNG CỘNG ngày bằng tổng dòng con, STT đếm lại; khối toàn kỳ không
    thêm truy vấn; ngày tách trang ghi "(tiếp)" và lặp TỔNG CỘNG; Excel hai sheet cùng khối"""
    A, B = van_don["A"], van_don["B"]
    _bao_cao(bang_mkt, A, "2026-08-01", "SP1", mess=10, don=2)
    _bao_cao(bang_mkt, B, "2026-08-01", "SP2", mess=30, don=3)
    _bao_cao(bang_mkt, A, "2026-08-02", "SP1", mess=10, don=1)
    client.force_login(B)
    query = {"nguon": bang_mkt.code, "tu": "2026-08-01", "den": "2026-08-02"}
    r = client.get("/bao-cao/tong-hop/", query)
    assert r.status_code == 200
    blocks = r.context["blocks"]
    cot = [c.label for c in r.context["result"].columns]
    assert [b["kind"] for b in blocks] == ["period", "day", "day"]
    ky, ngay2, ngay1 = blocks
    # Khối toàn kỳ: hai người, sắp theo mã, cộng cả kỳ; cột định danh STT · Team · Nhân sự · Leader
    assert ky["title"].startswith("Toàn kỳ 01/08 – 02/08/2026") and ky["count"] == 2
    assert [c["code"] for c in ky["identity_columns"]] == ["stt", "team", "person", "leader"] and ky["label_span"] == 4
    nguoi = {row["person"]: row for row in ky["rows"]}
    assert [row["stt"] for row in ky["rows"]] == [1, 2] and set(nguoi) == {employee_code(A), employee_code(B)}
    assert _o(nguoi[employee_code(A)], cot, "Số Mess") == "20" and _o(nguoi[employee_code(B)], cot, "Số Mess") == "30"
    assert _o(nguoi[employee_code(A)], cot, "Số đơn") == "3"
    assert nguoi[employee_code(A)]["team"] == (A.profile.team.name if A.profile.team_id else "Chưa có team")
    assert ky["total_label"] == "TỔNG CỘNG · toàn kỳ" and ky["totals"] == aggregations.total_cells(r.context["result"])
    # Khối ngày: mới nhất trước, không cột Ngày, TỔNG CỘNG ngày = tổng dòng con, STT từ 1
    assert ngay2["title"] == "02.08.2026" and ngay1["title"] == "01.08.2026"
    assert all("nhom" not in [c["code"] for c in b["identity_columns"]] for b in (ngay1, ngay2))
    assert [row["stt"] for row in ngay1["rows"]] == [1, 2] and [row["stt"] for row in ngay2["rows"]] == [1]
    assert dict(zip(cot, ngay1["totals"]))["Số Mess"] == "40" and dict(zip(cot, ngay1["totals"]))["Số đơn"] == "5"
    assert ngay1["total_label"] == "TỔNG CỘNG"
    html = r.content.decode()
    assert html.count('<table class="bang report-table"') == 3 and html.count('class="report-block-title"') == 3
    assert '<h3>02.08.2026</h3>' in html and 'data-pos="1" colspan="4">TỔNG CỘNG · toàn kỳ</th>' in html
    # Excel: sheet toàn kỳ và sheet theo ngày, cùng số
    book = load_workbook(BytesIO(client.get("/bao-cao/tong-hop/xuat/", query).content), data_only=True)
    ky_xls = list(book["Toan ky theo nhan su"].values)
    assert ky_xls[4][:4] == ("STT", "Team", "Nhân sự", "Leader") and str(ky_xls[5][0]).startswith("TỔNG CỘNG")
    assert {d[2] for d in ky_xls[6:8]} == {employee_code(A), employee_code(B)}
    dong_a = next(d for d in ky_xls[6:8] if d[2] == employee_code(A))
    assert dong_a[4 + cot.index("Số Mess")] == 20
    ngay_xls = [d[0] for d in book["Theo ngay"].values if d and d[0] is not None]
    assert ngay_xls == ["Ngày 02.08.2026", "STT", "TỔNG CỘNG", 1, "Ngày 01.08.2026", "STT", "TỔNG CỘNG", 1, 2]
    # Ngày bị tách trang: 13 ngày × 2 người = 26 dòng, trang 25 dòng → trang 2 còn một người của ngày 01.08
    for i in range(1, 13):
        ngay = (date(2026, 8, 1) + timedelta(days=i)).isoformat()
        _bao_cao(bang_mkt, A, ngay, "SP1", mess=10)
        _bao_cao(bang_mkt, B, ngay, "SP2", mess=30)
    query = {"nguon": bang_mkt.code, "tu": "2026-08-01", "den": "2026-08-13", "moi_trang": "25", "trang": "2"}
    r2 = client.get("/bao-cao/tong-hop/", query)
    khoi = r2.context["blocks"]
    assert khoi[0]["kind"] == "period" and khoi[0]["count"] == 2      # khối toàn kỳ vẫn đủ cả kỳ trên trang 2
    assert [b["title"] for b in khoi[1:]] == ["01.08.2026 (tiếp)"]
    assert len(khoi[1]["rows"]) == 1 and khoi[1]["rows"][0]["stt"] == 2
    assert dict(zip(cot, khoi[1]["totals"]))["Số Mess"] == "40"       # TỔNG CỘNG đủ cả ngày, không chỉ trang


def test_gop_chi_con_dong_tong_ngay(client, bang_mkt, mkt_source, van_don, nguoi_dung):
    """AC-42.7 — Gộp (`gop=1`): mỗi ngày một dòng là TỔNG CỘNG của ngày, phân trang theo ngày, khối
    toàn kỳ vẫn đứng đầu; chip Gộp có × về Không gộp; Excel sheet Theo ngay chỉ dòng ngày, cùng số"""
    A, B = van_don["A"], van_don["B"]
    _bao_cao(bang_mkt, A, "2026-08-01", "SP1", mess=10)
    _bao_cao(bang_mkt, B, "2026-08-01", "SP2", mess=30)
    _bao_cao(bang_mkt, A, "2026-08-02", "SP1", mess=10)
    client.force_login(B)
    query = {"nguon": bang_mkt.code, "tu": "2026-08-01", "den": "2026-08-02", "gop": "1"}
    r = client.get("/bao-cao/tong-hop/", query)
    blocks = r.context["blocks"]
    cot = [c.label for c in r.context["result"].columns]
    assert [b["kind"] for b in blocks] == ["period", "days"] and blocks[0]["count"] == 2
    ngay = blocks[1]
    assert [row["nhom"] for row in ngay["rows"]] == ["02.08.2026", "01.08.2026"]
    assert [_o(row, cot, "Số Mess") for row in ngay["rows"]] == ["10", "40"]
    assert [c["code"] for c in ngay["identity_columns"]] == ["nhom"] and r.context["ten_don_vi"] == "ngày"
    chips = {c["label"]: c for c in r.context["chips"]}
    assert chips["Gộp"]["url"] and "gop=" not in chips["Gộp"]["url"]
    html = r.content.decode()
    assert 'aria-pressed="true">Gộp</a>' in html and html.count('<table class="bang report-table"') == 2
    # Không gộp là mặc định: không có chip, hai khối ngày
    r0 = client.get("/bao-cao/tong-hop/", {k: v for k, v in query.items() if k != "gop"})
    assert "Gộp" not in {c["label"] for c in r0.context["chips"]} and [b["kind"] for b in r0.context["blocks"]] == ["period", "day", "day"]
    # Excel khi Gộp: sheet Theo ngay = hàng tiêu đề (Ngày, …), TỔNG CỘNG toàn kỳ, rồi mỗi ngày một dòng
    ngay_xls = list(load_workbook(BytesIO(client.get("/bao-cao/tong-hop/xuat/", query).content), data_only=True)["Theo ngay"].values)
    assert ngay_xls[0][0] == "Ngày" and str(ngay_xls[1][0]).startswith("TỔNG CỘNG")
    assert [d[0] for d in ngay_xls[2:4]] == ["02.08.2026", "01.08.2026"] and ngay_xls[3][1 + cot.index("Số Mess")] == 40
