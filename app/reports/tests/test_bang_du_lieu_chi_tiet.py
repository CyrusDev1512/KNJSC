"""ADR-042 đợt 4 — Bảng dữ liệu của bảng có nguồn báo cáo hiện thành báo cáo chi tiết theo ngày."""
import re
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from core.identity import employee_code
from reports.tests.test_aggregations import bang_mkt  # noqa: F401
from reports.tests.test_mkt_derived_revenue import _bao_cao, mkt_source, van_don  # noqa: F401

pytestmark = pytest.mark.django_db

CAD = Decimal("17500")


def _theo_nhan(result, raw):
    return dict(zip([c.label for c in result.columns], raw))


def test_bang_du_lieu_nguon_bao_cao_hien_chi_tiet_theo_ngay(client, bang_mkt, mkt_source, van_don, nguoi_dung,
                                                             django_assert_max_num_queries):
    """AC-42.13 — Bảng có nguồn báo cáo MKT mở ở Bảng dữ liệu là báo cáo chi tiết theo ngày dùng chung
    động cơ: mỗi lần nộp một dòng (hai lần nộp cùng ngày cùng người là hai dòng, STT riêng), khối toàn
    kỳ theo nhân sự đứng đầu, TỔNG CỘNG ngày đúng; (TT) trên dòng người chỉ khi cặp (ngày, người) nộp
    một lần, nộp nhiều lần thì "—" mà TỔNG CỘNG không cộng đôi; Gộp; `?dang=tho` về liệt kê thô có
    liên kết quay lại; bảng không có nguồn giữ nguyên; Staff chỉ thấy dòng mình, bộ phận khác 404;
    Excel cùng khối; form Ngưỡng màu cho quản lý; không quá 10 truy vấn"""
    A, B = van_don["A"], van_don["B"]
    _bao_cao(bang_mkt, A, "2026-08-01", "SP1", mess=10, don=2, cpqc="3")
    _bao_cao(bang_mkt, A, "2026-08-01", "SP2", mess=20, don=4, cpqc="5")   # A nộp lần hai cùng ngày
    _bao_cao(bang_mkt, B, "2026-08-01", "SP1", mess=5, don=1, cpqc="1")
    _bao_cao(bang_mkt, A, "2026-08-02", "SP1", mess=8, don=1, cpqc="2")
    url = f"/bang/{bang_mkt.code}/"
    ky = {"tu": "2026-08-01", "den": "2026-08-02"}

    client.force_login(B)                                  # manager_mkt: thấy cả bộ phận
    r = client.get(url, ky)
    assert r.status_code == 200 and r.context["khoi"] is True
    result = r.context["result"]
    blocks = r.context["blocks"]
    assert [b["kind"] for b in blocks] == ["period", "day", "day"]
    toan_ky, d02, d01 = blocks                              # ngày mới nhất trước
    assert d01["title"] == "01.08.2026" and d02["title"] == "02.08.2026"
    # Ngày 01.08: ba lần nộp là ba dòng (sắp theo mã rồi thứ tự nộp) — A hai dòng riêng, STT 1, 2, 3;
    # TỔNG CỘNG = 35 mess, 7 đơn
    assert [(row["person"], row["stt"]) for row in d01["rows"]] == [
        (employee_code(B), 1), (employee_code(A), 2), (employee_code(A), 3)]
    assert _theo_nhan(result, d01["totals_raw"])["Số Mess"] == 35
    assert _theo_nhan(result, d01["totals_raw"])["Số đơn"] == 7
    assert [_theo_nhan(result, row["raw"])["Số Mess"] for row in d01["rows"]] == [5, 10, 20]
    # (TT) khoá theo (ngày, người): A nộp hai lần ngày 01.08 → hai dòng đó "—", B một lần → có số;
    # TỔNG CỘNG ngày = w1 + w2 (A) + w4 (B) = 3 đơn, 300 CAD quy ₫ — không cộng đôi phần của A
    assert result.derived_shared == frozenset({(date(2026, 8, 1), employee_code(A))})
    tt = [(_theo_nhan(result, row["raw"])["Số đơn (TT)"], _theo_nhan(result, row["raw"])["DS Chốt (TT)"]) for row in d01["rows"]]
    assert tt == [(1, 200 * CAD), (None, None), (None, None)]
    assert str(d01["rows"][1]["cells"][[c.label for c in result.columns].index("Số đơn (TT)")]) == "—"
    assert _theo_nhan(result, d01["totals_raw"])["Số đơn (TT)"] == 3
    assert _theo_nhan(result, d01["totals_raw"])["DS Chốt (TT)"] == 300 * CAD
    assert _theo_nhan(result, d02["totals_raw"])["Số đơn (TT)"] == 1
    # Khối toàn kỳ: mỗi người một dòng cộng cả kỳ; A = 38 mess, 3 đơn (TT); B = 5 mess, 1 đơn (TT)
    assert [row["person"] for row in toan_ky["rows"]] == sorted([employee_code(A), employee_code(B)])
    theo_nguoi = {row["person"]: _theo_nhan(result, row["raw"]) for row in toan_ky["rows"]}
    assert theo_nguoi[employee_code(A)]["Số Mess"] == 38 and theo_nguoi[employee_code(A)]["Số đơn (TT)"] == 3
    assert theo_nguoi[employee_code(B)]["Số Mess"] == 5 and theo_nguoi[employee_code(B)]["Số đơn (TT)"] == 1
    assert _theo_nhan(result, toan_ky["totals_raw"])["Số Mess"] == 43
    assert _theo_nhan(result, toan_ky["totals_raw"])["Số đơn (TT)"] == 4
    assert _theo_nhan(result, toan_ky["totals_raw"])["DS Chốt (TT)"] == 325 * CAD
    html = r.content.decode()
    assert 'id="report-nguong"' in html and "Xem từng dòng thô" in html and 'rel="noopener">Mở trong KN CRM</a>' in html
    assert '<table class="bang bang-luoi">' not in html and "nộp nhiều lần thì (TT) chỉ hiện ở dòng TỔNG CỘNG" in html
    assert r.context["moi_trang"] == 25                     # quy tắc 1: mặc định 25 dòng
    # Gộp: mỗi ngày một dòng
    r_gop = client.get(url, {**ky, "gop": "1"})
    assert [b["kind"] for b in r_gop.context["blocks"]] == ["period", "days"]
    assert [row["nhom"] for row in r_gop.context["blocks"][1]["rows"]] == ["02.08.2026", "01.08.2026"]
    # Liệt kê thô: bảng cũ, có liên kết quay lại, phân trang và Xoá lọc giữ `dang=tho`
    r_tho = client.get(url, {"dang": "tho"})
    html_tho = r_tho.content.decode()
    assert r_tho.status_code == 200 and not r_tho.context.get("khoi")
    assert '<table class="bang bang-luoi">' in html_tho and "Xem báo cáo theo ngày" in html_tho
    assert r_tho.context["qs_loc"] == "&dang=tho" and 'href="/bang/' + bang_mkt.code + '/?dang=tho">Xoá lọc</a>' in html_tho
    assert r_tho.context["page_obj"].paginator.count == 4
    # Excel: hai sheet như Báo cáo tổng hợp, tổng toàn kỳ 43 mess
    r_xls = client.get(reverse("bang_xuat", args=[bang_mkt.code]), ky)
    assert r_xls.status_code == 200 and r_xls["Content-Type"].endswith("sheet")
    book = load_workbook(BytesIO(r_xls.content))
    assert "Toan ky theo nhan su" in book.sheetnames and "Theo ngay" in book.sheetnames
    sheet = book["Toan ky theo nhan su"]
    tong = next(row for row in sheet.iter_rows(values_only=True) if row and row[0] and "TỔNG CỘNG" in str(row[0]))
    assert 43 in tong
    # Ngân sách truy vấn (Q2)
    client.get(url, ky)
    with django_assert_max_num_queries(10):
        assert client.get(url, ky).status_code == 200
    # Staff chỉ thấy dòng của mình, không thấy form Ngưỡng màu
    client.force_login(A)
    r_a = client.get(url, ky)
    assert [row["person"] for row in r_a.context["blocks"][0]["rows"]] == [employee_code(A)]
    assert employee_code(B) not in r_a.content.decode() and 'id="report-nguong"' not in r_a.content.decode()
    # Bộ phận khác: 404 như mọi bảng
    client.force_login(nguoi_dung["manager_sale"])
    assert client.get(url, ky).status_code == 404
    # Bảng không có nguồn báo cáo (vận đơn) giữ nguyên liệt kê thô
    client.force_login(nguoi_dung["admin"])
    r_vd = client.get(f"/bang/{van_don['table'].code}/")
    assert r_vd.status_code == 200 and not r_vd.context.get("khoi") and '<table class="bang bang-luoi">' in r_vd.content.decode()
