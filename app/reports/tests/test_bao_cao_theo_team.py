"""Cách xem "Theo team" thay "Hiệu suất theo phòng ban" trên Báo cáo tổng hợp — AC-22.20 (chủ dự án 30.09.2026).

Bảng báo cáo nào cũng thuộc đúng một bộ phận nên cách xem phòng ban chỉ ra một dòng = TỔNG CỘNG. Theo team:
mỗi team một dòng (mỗi loại tiền một dòng, ADR-046), kèm Leader của team; dòng chưa có team gom "Chưa có team";
DS Chốt (TT) đối soát theo team của marketer phụ trách vận đơn; Leader chỉ thấy team mình.
"""
from datetime import date

import pytest

from core.constants import Rank
from core.identity import employee_code
from org.models import Team
from reports import aggregations
from reports.services import activity_service
from reports.tests.test_aggregations import bang_mkt  # noqa: F401
from reports.tests.test_activity import delivery_source, marketing_scope  # noqa: F401
from reports.tests.test_mkt_derived_revenue import _bao_cao, mkt_source, van_don  # noqa: F401

pytestmark = pytest.mark.django_db
KY = dict(start=date(2026, 8, 1), end=date(2026, 8, 2))


@pytest.fixture
def hai_team(nguoi_dung, departments, make_user):
    t1 = Team.objects.create(name="MKT 1", department=departments["mkt"])
    t2 = Team.objects.create(name="MKT 2", department=departments["mkt"])
    leader = make_user("leader_mkt_1", Rank.LEADER, departments["mkt"], t1)
    t1.leader = leader
    t1.save(update_fields=["leader"])
    a = nguoi_dung["staff_mkt"]
    a.profile.team = t1
    a.profile.save(update_fields=["team"])
    c = make_user("staff_mkt_2", Rank.STAFF, departments["mkt"], t2)
    return {"t1": t1, "t2": t2, "leader": leader, "A": a, "C": c}


def _dong(result):
    out = {}
    for item in result.rows:
        nhom, raw = aggregations.row_values(item, result)
        out[aggregations.format_group(nhom, result)] = (item, dict(zip([c.label for c in result.columns], raw)))
    return out


def test_cach_xem_theo_team_thay_phong_ban():
    """AC-22.20 — Danh sách Cách xem có "Hiệu suất theo team", không còn "Hiệu suất theo phòng ban\""""
    nhan = dict(activity_service.GROUPS)
    assert nhan.get("team") == "Hiệu suất theo team" and "department" not in nhan


def test_moi_team_mot_dong_kem_leader_va_doi_soat(bang_mkt, mkt_source, van_don, hai_team, nguoi_dung):  # noqa: F811
    """AC-22.20 — Theo team: mỗi team một dòng cộng đúng các lần nộp của người trong team, cột Leader là mã
    leader của team; người chưa có team gom "Chưa có team"; DS Chốt (TT) theo team của marketer phụ trách"""
    A, C, B = hai_team["A"], hai_team["C"], van_don["B"]
    _bao_cao(bang_mkt, A, "2026-08-01", "SP1", mess=10)
    _bao_cao(bang_mkt, A, "2026-08-02", "SP1", mess=5)
    _bao_cao(bang_mkt, C, "2026-08-01", "SP2", mess=7)
    _bao_cao(bang_mkt, B, "2026-08-01", "SP2", mess=3)
    result = activity_service.build(B, mkt_source, group="team", **KY)
    assert result.ok and result.group_label == "Team" and result.show_leader
    rows = _dong(result)
    assert set(rows) == {"MKT 1", "MKT 2", "Chưa có team"}
    assert rows["MKT 1"][1]["Số Mess"] == 15 and rows["MKT 2"][1]["Số Mess"] == 7 and rows["Chưa có team"][1]["Số Mess"] == 3
    assert rows["MKT 1"][0]["leader_name"] == employee_code(hai_team["leader"])
    # Vận đơn: A (MKT 1) phụ trách 60 + 40 + 25; B (chưa có team) 200
    assert rows["MKT 1"][1]["DS Chốt (TT)"] == 125 and rows["Chưa có team"][1]["DS Chốt (TT)"] == 200


def test_leader_chi_thay_team_minh(bang_mkt, mkt_source, van_don, hai_team):  # noqa: F811
    """AC-22.20 — Leader MKT 1 xem Theo team chỉ thấy dòng MKT 1 (phạm vi quyền giữ nguyên)"""
    _bao_cao(bang_mkt, hai_team["A"], "2026-08-01", "SP1")
    _bao_cao(bang_mkt, hai_team["C"], "2026-08-01", "SP2")
    rows = _dong(activity_service.build(hai_team["leader"], mkt_source, group="team", **KY))
    assert set(rows) == {"MKT 1"}


def test_man_hinh_theo_team_va_url_cu_phong_ban(client, bang_mkt, mkt_source, van_don, hai_team, nguoi_dung):  # noqa: F811
    """AC-22.20 — Màn Báo cáo tổng hợp không còn ô Cách xem (01.10.2026, AC-22.23): đường dẫn cũ `nhom=team`,
    `nhom=department` (đánh dấu trang, tệp đã gửi) vẫn mở, ra từng lần nộp có cột Team và Leader; Excel cũng có
    cột Team, Leader và dòng của team"""
    _bao_cao(bang_mkt, hai_team["A"], "2026-08-01", "SP1")
    client.force_login(nguoi_dung["admin"])
    for nhom in ("team", "department"):
        r = client.get("/bao-cao/tong-hop/", {"nguon": mkt_source.table.code, "nhom": nhom, "tu": "2026-08-01", "den": "2026-08-02"})
        html = r.content.decode()
        assert r.status_code == 200, nhom
        assert r.context["params"]["group"] == "day" and 'name="nhom"' not in html, nhom
        assert "MKT 1" in html and "Phòng ban" not in html
    # Tải Excel theo team: có cột Team và Leader, dòng MKT 1
    from io import BytesIO
    from openpyxl import load_workbook
    x = client.get("/bao-cao/tong-hop/xuat/", {"nguon": mkt_source.table.code, "nhom": "team", "tu": "2026-08-01",
                                              "den": "2026-08-02"})
    assert x.status_code == 200
    o = [[c for c in row] for row in load_workbook(BytesIO(x.content)).active.iter_rows(values_only=True)]
    tieu_de = next(row for row in o if row and "Team" in row)
    assert "Leader" in tieu_de and any(row and "MKT 1" in row for row in o)


def test_nguon_van_don_theo_team(client, delivery_source, nguoi_dung, teams):  # noqa: F811
    """AC-22.20 — Nguồn Vận đơn xem Theo team: team của người phụ trách Vận đơn, kèm Leader; dòng chưa phân
    công gom "Chưa có team"; màn hình mở được (không lỗi thiếu cột Leader)"""
    result = activity_service.build(nguoi_dung["admin"], delivery_source, group="team")
    rows = {r["nhom"]: r for r in result.rows}
    assert set(rows) == {teams["sale1"].name, teams["sale2"].name, "Chưa có team"}
    assert rows[teams["sale1"].name]["c_orders"] == 2 and rows[teams["sale2"].name]["c_orders"] == 1
    assert rows[teams["sale1"].name]["leader_name"] == employee_code(teams["leader1"])
    client.force_login(nguoi_dung["admin"])
    r = client.get("/bao-cao/tong-hop/", {"nguon": delivery_source.table.code, "nhom": "team"})
    assert r.status_code == 200 and teams["sale1"].name in r.content.decode()
