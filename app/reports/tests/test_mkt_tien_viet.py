"""Báo cáo Marketing nộp bằng tiền Việt — ADR-047 (chủ dự án 03.10.2026).

"Tất cả số trong báo cáo được nộp đều là tiền Việt": Loại tiền của báo cáo Marketing luôn là VND, không theo
Thị trường; báo cáo cũ đổi nhãn sang VND, số giữ nguyên; không có tỉ giá ở đâu cả. Báo cáo Sale giữ như cũ.
"""
import importlib
from datetime import date

import pytest
from django.apps import apps as django_apps

from core.identity import employee_code
from forms_builder.models import DataRecord
from forms_builder.services import record_service
from orders.services import currency_service
from reports import aggregations
from reports.services import activity_service, daily_service
from reports.tests.test_aggregations import bang_mkt  # noqa: F401
from reports.tests.test_mkt_derived_revenue import _bao_cao, mkt_source, van_don  # noqa: F401

pytestmark = pytest.mark.django_db


def _tien(row):
    row.refresh_from_db()
    return row.data.get("loai_tien")


def test_loai_tien_bao_cao_khai_mot_cho():
    """AC-47.1 — Một chỗ quyết định loại tiền của báo cáo ngày: Marketing luôn VND dù Thị trường nào (kể cả
    trống); Sale theo Thị trường như ADR-031"""
    assert currency_service.report_currency("mkt", "Canada") == "VND"
    assert currency_service.report_currency("mkt", "Hoa Kỳ") == "VND"
    assert currency_service.report_currency("mkt", "", allow_empty=True) == "VND"
    assert currency_service.report_currency("sale", "Canada") == "CAD"
    assert currency_service.report_currency("sale", "", allow_empty=True) == ""


def test_nop_va_sua_bao_cao_mkt_luon_vnd(client, bang_mkt, mkt_source, nguoi_dung):  # noqa: F811
    """AC-47.2 — Nộp báo cáo MKT chọn Canada → lưu VND; đổi Thị trường trên dòng (sửa ô) vẫn VND; form hiện
    Loại tiền VND với mọi thị trường và ô tiền ghi "(₫)" """
    A = nguoi_dung["staff_mkt"]
    dong = _bao_cao(bang_mkt, A, "2026-08-01", "SP1")          # Thị trường Canada
    assert _tien(dong) == "VND"
    record_service.update_cell(dong, "thi_truong", "Hoa Kỳ", actor=nguoi_dung["admin"])
    assert _tien(dong) == "VND"
    gia_tri = daily_service.protected_values(mkt_source.table.forms.get(), {"thi_truong": "Canada"},
                                             list(mkt_source.table.forms.get().ordered_fields()),
                                             date(2026, 8, 1), A)
    assert [v for k, v in gia_tri.items() if k.endswith("loai_tien")] == ["VND"]
    form = mkt_source.table.forms.get()
    client.force_login(A)
    html = client.get(f"/bao-cao/?bieu_mau={form.code}").content.decode()
    assert '"Canada": "VND"' in html and '"Hoa K\\u1ef3": "VND"' in html
    assert "CPQC (₫)" in html and "Doanh số (₫)" in html and "Số Mess (₫)" not in html


def test_bao_cao_sale_van_theo_thi_truong(bang_mkt, mkt_source, nguoi_dung):  # noqa: F811
    """AC-47.2 — Báo cáo Sale không đổi: Canada vẫn CAD"""
    from reports.models import ReportSource
    from forms_builder.models import TableDef
    ReportSource.objects.filter(pk=mkt_source.pk).update(kind="sale")
    dong = _bao_cao(TableDef.objects.get(pk=bang_mkt.pk), nguoi_dung["staff_mkt"], "2026-08-01", "SP1")
    assert _tien(dong) == "CAD"


def test_doi_nhan_bao_cao_cu_sang_vnd_va_chay_nguoc(bang_mkt, mkt_source, nguoi_dung):  # noqa: F811
    """AC-47.3 — Tệp chuyển đổi: dòng báo cáo MKT cũ mang CAD/USD/trống đổi nhãn sang VND, số giữ nguyên; bảng
    Sale không đổi; chạy ngược thì nhãn lấy lại theo Thị trường (Canada → CAD, trống → trống)"""
    from reports.models import ReportSource
    A = nguoi_dung["staff_mkt"]
    cad = _bao_cao(bang_mkt, A, "2026-08-01", "SP1", cpqc="13250000")
    my = _bao_cao(bang_mkt, A, "2026-08-02", "SP1")
    trong = _bao_cao(bang_mkt, A, "2026-08-03", "SP1")
    DataRecord.all_objects.filter(pk=cad.pk).update(data={**cad.data, "loai_tien": "CAD"})
    DataRecord.all_objects.filter(pk=my.pk).update(data={**my.data, "thi_truong": "Hoa Kỳ", "loai_tien": "USD"})
    DataRecord.all_objects.filter(pk=trong.pk).update(data={**trong.data, "thi_truong": "", "loai_tien": ""})
    migration = importlib.import_module("reports.migrations.0006_bao_cao_mkt_tien_viet")
    migration.doi_sang_vnd(django_apps, None)
    assert [_tien(x) for x in (cad, my, trong)] == ["VND", "VND", "VND"]
    assert str(cad.data["cpqc"]).startswith("13250000")
    migration.tra_lai(django_apps, None)
    assert [_tien(x) for x in (cad, my, trong)] == ["CAD", "USD", ""]
    # Bảng Sale không bị đụng
    ReportSource.objects.filter(pk=mkt_source.pk).update(kind="sale")
    migration.doi_sang_vnd(django_apps, None)
    assert _tien(cad) == "CAD"


def test_bao_cao_tong_hop_mkt_mot_tong_vnd_va_so_don_tt(bang_mkt, mkt_source, van_don, nguoi_dung):  # noqa: F811
    """AC-47.4 — Báo cáo tổng hợp MKT: một dòng TỔNG CỘNG · VND; Số đơn (TT) vẫn đếm đúng đơn của marketer dù
    đơn bằng CAD/USD; DS Chốt (TT) để trống vì không quy đổi tỉ giá; số tiền đúng như nhập"""
    A, B = van_don["A"], van_don["B"]
    _bao_cao(bang_mkt, A, "2026-08-01", "SP1", mess=10, cpqc="13250000")
    _bao_cao(bang_mkt, B, "2026-08-01", "SP2", mess=20)
    result = activity_service.build(B, mkt_source, group="day", start=date(2026, 8, 1), end=date(2026, 8, 1))
    assert result.ok
    tong = aggregations.total_rows(result)
    assert [tien for tien, _ in tong] == ["VND"]
    nhan = [c.label for c in result.columns]
    dong = {}
    for item in result.rows:
        _, raw = aggregations.row_values(item, result)
        dong[item["person_name"]] = dict(zip(nhan, raw))
    assert dong[employee_code(A)]["Số đơn (TT)"] == 2 and dong[employee_code(B)]["Số đơn (TT)"] == 1
    assert dong[employee_code(A)]["DS Chốt (TT)"] is None
    assert dong[employee_code(A)]["CPQC"] == 13250000


def test_bao_cao_mkt_khong_con_cot_hoa_don(client, bang_mkt, mkt_source, nguoi_dung):  # noqa: F811
    """AC-47.5 — Báo cáo tổng hợp MKT (màn hình, Excel) không còn cột Hóa đơn và Hóa đơn/DS Chốt (TT); dữ liệu cột
    Hóa đơn trong bảng giữ nguyên; các cột khác giữ thứ tự"""
    from io import BytesIO
    from openpyxl import load_workbook
    dong = _bao_cao(bang_mkt, nguoi_dung["staff_mkt"], "2026-08-01", "SP1", hoa_don="8")
    result = activity_service.build(nguoi_dung["manager_mkt"], mkt_source, start=date(2026, 8, 1), end=date(2026, 8, 1))
    nhan = [c.label for c in result.columns]
    assert "Hóa đơn" not in nhan and "Hóa đơn/DS Chốt (TT)" not in nhan
    assert nhan[:4] == ["Số Mess", "CPQC", "Số đơn", "Số đơn (TT)"] and "DS Chốt (TT)" in nhan and "AOV" in nhan
    dong.refresh_from_db()
    assert str(dong.data["hoa_don"]) == "8"
    client.force_login(nguoi_dung["manager_mkt"])
    ky = {"nguon": bang_mkt.code, "tu": "2026-08-01", "den": "2026-08-01"}
    html = client.get("/bao-cao/tong-hop/", ky).content.decode()
    assert ">Hóa đơn<" not in html and "Hóa đơn/DS Chốt" not in html
    sach = load_workbook(BytesIO(client.get("/bao-cao/tong-hop/xuat/", ky).content), data_only=True)
    chu = " ".join(str(o) for ws in sach.worksheets for hang in ws.iter_rows(values_only=True) for o in hang if o)
    assert "Hóa đơn" not in chu
