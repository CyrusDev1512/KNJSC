"""Báo cáo Sale ẩn cột Lần nộp — ADR-047 bổ sung 04.10.2026.

Chủ dự án thử Staging 04.10.2026: "cái lần nộp này đã bỏ rồi mà" — cột Lần nộp mới chỉ ẩn ở báo cáo MKT (AC-47.6).
Báo cáo Sale cũng ẩn cột Lần nộp ở màn hình, Bảng dữ liệu dạng báo cáo và Excel; khác MKT, Sale **giữ cột Loại tiền**
vì một người nộp được nhiều loại tiền và không cộng lẫn tiền tệ (ADR-046). Chỉ đổi hiển thị, không đổi cách cộng tổng.
"""
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from forms_builder.models import TableDef
from forms_builder.services import record_service
from reports.models import ReportSource
from reports.tests.test_aggregations import bang_mkt  # noqa: F401
from reports.tests.test_mkt_derived_revenue import mkt_source  # noqa: F401

pytestmark = pytest.mark.django_db


def _nop(bang, actor, ngay, thi_truong, mess):
    values = {"ngay": ngay, "marketer": "x", "san_pham": "SP1", "so_mess": mess, "cpqc": "3", "so_don": 2,
              "doanh_so": "100", "thi_truong": thi_truong}
    return record_service.create_record(bang, values, actor=actor, system_day=date.fromisoformat(ngay))


def test_bao_cao_sale_an_lan_nop_giu_loai_tien(client, bang_mkt, mkt_source, nguoi_dung):  # noqa: F811
    """AC-47.7 — Báo cáo tổng hợp, Bảng dữ liệu dạng báo cáo và Excel của nguồn Sale không còn cột Lần nộp ở khối
    toàn kỳ, khối ngày và khối Gộp nhưng giữ cột Loại tiền; hai lần nộp cùng ngày vẫn là hai dòng; TỔNG CỘNG vẫn
    tách theo loại tiền"""
    ReportSource.objects.filter(pk=mkt_source.pk).update(kind="sale")     # Sale: Loại tiền theo Thị trường
    bang = TableDef.objects.get(pk=bang_mkt.pk)                            # bản mới: không giữ nguồn MKT đã nạp sẵn
    A = nguoi_dung["staff_mkt"]
    _nop(bang, A, "2026-08-01", "Canada", 10)
    _nop(bang, A, "2026-08-01", "Canada", 20)
    _nop(bang, A, "2026-08-01", "Hoa Kỳ", 30)
    client.force_login(nguoi_dung["manager_mkt"])
    ky = {"tu": "2026-08-01", "den": "2026-08-01"}
    tong_hop = {"nguon": bang_mkt.code, **ky}
    for url, query in (("/bao-cao/tong-hop/", tong_hop), (f"/bang/{bang_mkt.code}/", ky)):
        for gop in ({}, {"gop": "1"}):
            r = client.get(url, {**query, **gop})
            assert r.status_code == 200
            for b in r.context["blocks"]:
                ma = [c["code"] for c in b["identity_columns"]]
                assert "lan" not in ma and "tien" in ma, (url, gop, b["kind"], ma)
            html = r.content.decode()
            assert "id-lan" not in html and "id-tien" in html, (url, gop)
    r = client.get("/bao-cao/tong-hop/", tong_hop)
    ngay = r.context["blocks"][1]
    assert ngay["kind"] == "day" and len(ngay["rows"]) == 3           # hai lần nộp Canada vẫn hai dòng, thêm dòng USD
    assert sorted(t["currency"] for t in ngay["total_rows"]) == ["CAD", "USD"]
    sach = load_workbook(BytesIO(client.get("/bao-cao/tong-hop/xuat/", tong_hop).content), data_only=True)
    chu = {str(o) for ws in sach.worksheets for hang in ws.iter_rows(values_only=True) for o in hang if o}
    assert "Lần nộp" not in chu and "Loại tiền" in chu
