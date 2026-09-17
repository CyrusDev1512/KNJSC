"""Nhập/xuất bằng Celery thật trong DB test; một job mỗi loại mỗi vòng."""
import json
import os
import time
import uuid
from pathlib import Path
import django
django.setup()
from django.conf import settings
from django.db import connection
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.http import QueryDict
from core.constants import Rank,JobStatus
from core.models import BackgroundJob
from org.models import UserProfile
from forms_builder.models import TableDef,DataRecord
from forms_builder.services import import_service,export_service
from openpyxl import load_workbook

assert connection.settings_dict['HOST']=='knjsc-crm-update-db'
assert connection.settings_dict['NAME'].startswith('test_crm_update_load_')
assert 'knjsc-crm-update-redis' in settings.CELERY_BROKER_URL
assert not settings.CELERY_TASK_ALWAYS_EAGER
table=TableDef.objects.get(code='van_don')
user,_=get_user_model().objects.get_or_create(username='background_test')
UserProfile.objects.get_or_create(user=user,defaults={'rank':Rank.MANAGER,'department':table.department,'must_change_password':False})
evidence=[]
for iteration in range(int(os.environ.get('BACKGROUND_ROUNDS','1'))):
    token=uuid.uuid4().hex
    csv='ma_don,ten_khach,so_dien_thoai,so_don,cpqc\n'+''.join(f'BG-{token}-{i},Khach background,0012345678,5,100.00\n' for i in range(100))
    start=time.monotonic()
    job=import_service.prepare(table,SimpleUploadedFile('synthetic.csv',csv.encode()),actor=user)
    import_service.confirm(job,actor=user)
    jobs=[job]
    kind,export=export_service.export(user,table,QueryDict('f_ten_khach__chua=Khach+test+12'),builder='grid')
    assert kind=='job','Bộ lọc cần đủ dòng để xuất nền'
    jobs.append(export)
    for job in jobs:
        deadline=time.monotonic()+120
        while time.monotonic()<deadline:
            job.refresh_from_db()
            if job.status in (JobStatus.DONE,JobStatus.FAILED):break
            time.sleep(.2)
        assert job.status==JobStatus.DONE,(job.pk,job.status,job.error)
        evidence.append({'job':job.pk,'kind':job.kind,'total':job.total,'elapsed_ms':round((time.monotonic()-start)*1000,2),'status':job.status,
            'queued_ms':round((job.started_at-job.created_at).total_seconds()*1000,2),
            'run_ms':round((job.finished_at-job.started_at).total_seconds()*1000,2)})
    assert DataRecord.objects.filter(table=table,data__ma_don__startswith='BG-'+token).count()==100
    export_service.check_download(export,user)
    result_file=export_service.result_file(export)
    assert result_file.exists()
    workbook=load_workbook(result_file,read_only=True,data_only=True)
    try:
        rows=workbook.active.iter_rows(values_only=True)
        headers=next(rows)
        code_column=headers.index('ma_don')
        phone_column=headers.index('so_dien_thoai')
        codes=[]
        for row in rows:
            codes.append(row[code_column])
            assert isinstance(row[phone_column],str) and row[phone_column].startswith('0')
        expected=set(DataRecord.objects.filter(pk__in=export.summary['exported_row_ids']).values_list('data__ma_don',flat=True))
        assert len(codes)==len(set(codes))==export.total and set(codes)==expected
        evidence[-1]['verified_excel_rows']=len(codes)
    finally:
        workbook.close()
    Path('/runtime/background.json').write_text(json.dumps(evidence,indent=2))
    print(json.dumps(evidence[-2:]),flush=True)
    if iteration+1<int(os.environ.get('BACKGROUND_ROUNDS','1')):time.sleep(max(0,60-(time.monotonic()-start)))
